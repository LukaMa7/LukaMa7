# -*- coding: utf-8 -*-
"""Shared helpers for the ECC-2:2024 training attachments (RTL DOCX / XLSX).

All generated content is fictional training data for the fictional entity
«هيئة المدار للخدمات الرقمية». Regenerate everything with build_all.py.
"""
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SP = os.path.dirname(HERE)
OUT = os.path.join(SP, 'deliver', 'attachments')

FONT = 'Arial'
LRM = '‎'
ENTITY = 'هيئة المدار للخدمات الرقمية'
ENTITY_FULL = 'هيئة المدار للخدمات الرقمية (جهة حكومية افتراضية لأغراض التدريب)'
MARK = 'بيانات تدريبية افتراضية'
FOOTER = MARK + ' — ' + ENTITY
PERIOD_FROM = '2026-07-01'
PERIOD_TO = '2026-09-30'

HEAD_FILL = 'D9E2F3'   # header row shading
LABEL_FILL = 'EEF2F8'  # label column shading in key/value tables
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
RED = RGBColor(0xA6, 0x1B, 0x1B)
GREY = RGBColor(0x55, 0x55, 0x55)

# ---------------------------------------------------------------- bidi text
_NOT = r'A-Za-z0-9\-.:'   # '_' and '/' allowed around codes so folder names like 02_..._2-9/ stay LTR
_TIME_RANGE = r'\d{1,2}:\d{2}\s?-\s?\d{1,2}:\d{2}'
_DATE = r'\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}(?::\d{2})?)?'
_CODE = r'\d{1,2}(?:-\d{1,2}){1,3}'
_RX = re.compile(rf'(?<![{_NOT}])(?:{_TIME_RANGE}|{_DATE}|{_CODE})(?![A-Za-z0-9\-])')


def ltr(text):
    """Wrap control codes, dates and date-times with U+200E so they stay LTR."""
    if text is None:
        return ''
    text = str(text).replace(LRM, '')
    return _RX.sub(lambda m: LRM + m.group(0) + LRM, text)


# ---------------------------------------------------------------- oxml utils
_RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike',
              'outline', 'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish', 'webHidden',
              'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect',
              'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang', 'eastAsianLayout',
              'specVanish', 'oMath']
_PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr',
              'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
              'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
              'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc',
              'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle',
              'rPr', 'sectPr', 'pPrChange']
_TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
               'textDirection', 'tcFitText', 'vAlign', 'hideMark']
_TRPR_ORDER = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit',
               'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden']
_SECT_ORDER = ['headerReference', 'footerReference', 'footnotePr', 'endnotePr', 'type', 'pgSz', 'pgMar',
               'paperSrc', 'pgBorders', 'lnNumType', 'pgNumType', 'cols', 'formProt', 'vAlign',
               'noEndnote', 'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid', 'printerSettings']


def _set_child(parent, tag, order, attrs=None):
    """Get or insert child <w:tag> respecting schema order; set attributes."""
    el = parent.find(qn('w:' + tag))
    if el is None:
        el = OxmlElement('w:' + tag)
        idx = order.index(tag)
        later = set(order[idx + 1:])
        inserted = False
        for i, ch in enumerate(list(parent)):
            local = ch.tag.split('}')[-1]
            if local in later:
                ch.addprevious(el)
                inserted = True
                break
        if not inserted:
            parent.append(el)
    for k, v in (attrs or {}).items():
        el.set(qn('w:' + k), str(v))
    return el


def set_run_fonts(rpr, size=None, bold=None, color=None, italic=None):
    rf = _set_child(rpr, 'rFonts', _RPR_ORDER)
    for att in list(rf.attrib):
        del rf.attrib[att]
    for k in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rf.set(qn('w:' + k), FONT)
    if bold is not None:
        for t in ('b', 'bCs'):
            el = _set_child(rpr, t, _RPR_ORDER)
            el.set(qn('w:val'), '1' if bold else '0')
    if italic is not None:
        for t in ('i', 'iCs'):
            el = _set_child(rpr, t, _RPR_ORDER)
            el.set(qn('w:val'), '1' if italic else '0')
    if color is not None:
        _set_child(rpr, 'color', _RPR_ORDER, {'val': str(color)})
    if size is not None:
        hp = str(int(round(size * 2)))
        _set_child(rpr, 'sz', _RPR_ORDER, {'val': hp})
        _set_child(rpr, 'szCs', _RPR_ORDER, {'val': hp})


_NBH = re.compile(r'(?<=[A-Za-z0-9])-(?=[A-Za-z0-9])')


def nb_hyphens(run):
    """Turn hyphens inside codes, dates and IDs into <w:noBreakHyphen/>.

    NOT USED: rendering checks showed that the non-breaking hyphen (bidi class ON) reverses
    LRM-wrapped codes such as 4-1-2 in LibreOffice. Kept only as a record of that finding.
    """
    text = run.text
    if not _NBH.search(text):
        return run
    r = run._r
    for t in r.findall(qn('w:t')):
        r.remove(t)
    parts = _NBH.split(text)
    for i, part in enumerate(parts):
        if i:
            r.append(OxmlElement('w:noBreakHyphen'))
        if part:
            t = OxmlElement('w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = part
            r.append(t)
    return run


def set_par_bidi(p_el):
    ppr = p_el.get_or_add_pPr()
    _set_child(ppr, 'bidi', _PPR_ORDER)
    return ppr


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    _set_child(tcpr, 'shd', _TCPR_ORDER, {'val': 'clear', 'color': 'auto', 'fill': fill})


# ---------------------------------------------------------------- document
def _style_fix(doc):
    styles = doc.styles
    # document defaults
    dd = styles.element.find(qn('w:docDefaults'))
    rpr_def = dd.find(qn('w:rPrDefault'))
    rpr = rpr_def.find(qn('w:rPr'))
    set_run_fonts(rpr, size=11)
    _set_child(rpr, 'lang', _RPR_ORDER, {'val': 'en-US', 'eastAsia': 'en-US', 'bidi': 'ar-SA'})
    ppr_def = dd.find(qn('w:pPrDefault'))
    if ppr_def is None:
        ppr_def = OxmlElement('w:pPrDefault')
        dd.append(ppr_def)
    ppr = ppr_def.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        ppr_def.append(ppr)
    _set_child(ppr, 'bidi', _PPR_ORDER)

    normal = styles['Normal']
    set_run_fonts(normal.element.get_or_add_rPr(), size=11)
    nppr = normal.element.get_or_add_pPr()
    _set_child(nppr, 'bidi', _PPR_ORDER)
    _set_child(nppr, 'spacing', _PPR_ORDER, {'before': '0', 'after': '100', 'line': '276', 'lineRule': 'auto'})

    sizes = {1: 16, 2: 13.5, 3: 12}
    for lvl, sz in sizes.items():
        st = styles['Heading %d' % lvl]
        set_run_fonts(st.element.get_or_add_rPr(), size=sz, bold=True, color='1F3A5F', italic=False)
        hppr = st.element.get_or_add_pPr()
        _set_child(hppr, 'bidi', _PPR_ORDER)
        _set_child(hppr, 'spacing', _PPR_ORDER, {'before': '240' if lvl == 1 else '180', 'after': '100'})
        _set_child(hppr, 'keepNext', _PPR_ORDER)
    for name in ('Title', 'Subtitle', 'Table Grid', 'Header', 'Footer', 'List Paragraph'):
        try:
            st = styles[name]
        except KeyError:
            continue
        set_run_fonts(st.element.get_or_add_rPr())
        if st.type == 1:  # paragraph style
            _set_child(st.element.get_or_add_pPr(), 'bidi', _PPR_ORDER)


def _add_page_field(par, size=8.5):
    run = par.add_run()
    set_run_fonts(run._r.get_or_add_rPr(), size=size, color='555555')
    for kind, txt in (('begin', None), (None, ' PAGE '), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), kind)
            run._r.append(fc)
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = txt
            run._r.append(it)


class RDoc:
    """A4 RTL Word document with fixed footer and optional trainer-only banner."""

    def __init__(self, header_text, trainer_only=False, title=None):
        self.d = Document()
        _style_fix(self.d)
        self.trainer_only = trainer_only
        sec = self.d.sections[0]
        self._setup_section(sec, header_text)
        cp = self.d.core_properties
        cp.author = 'فريق التدريب — بيانات افتراضية'
        cp.title = title or header_text
        cp.comments = MARK
        cp.keywords = 'ECC-2:2024; ' + MARK

    def _setup_section(self, sec, header_text, landscape=False):
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.header_distance = sec.footer_distance = Cm(1.0)
        _set_child(sec._sectPr, 'bidi', _SECT_ORDER)
        # header
        hp = sec.header.paragraphs[0]
        hp.text = ''
        set_par_bidi(hp._p)
        r = hp.add_run(ltr(header_text))
        set_run_fonts(r._r.get_or_add_rPr(), size=8.5, color='555555')
        if self.trainer_only:
            hp2 = sec.header.add_paragraph()
            set_par_bidi(hp2._p)
            r2 = hp2.add_run('للمدرب فقط — لا يوزَّع على المتدربين')
            set_run_fonts(r2._r.get_or_add_rPr(), size=8.5, bold=True, color='A61B1B')
        # footer
        fp = sec.footer.paragraphs[0]
        fp.text = ''
        set_par_bidi(fp._p)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(FOOTER + ' — صفحة ')
        set_run_fonts(r._r.get_or_add_rPr(), size=8.5, color='555555')
        _add_page_field(fp)

    # ------------------------------------------------------------ text
    def _runs(self, par, text, size=None, bold=False, color=None, italic=False):
        """Add runs; **x** = bold segment, [[x]] = coloured label segment."""
        text = ltr(text)
        tokens = re.split(r'(\*\*.+?\*\*|\[\[.+?\]\])', text)
        for tok in tokens:
            if not tok:
                continue
            b, c = bold, color
            if tok.startswith('**') and tok.endswith('**'):
                tok, b = tok[2:-2], True
            elif tok.startswith('[[') and tok.endswith(']]'):
                tok, b, c = '«' + tok[2:-2] + '»', True, '1F6F5C'
            r = par.add_run(tok)
            set_run_fonts(r._r.get_or_add_rPr(), size=size, bold=b, color=c, italic=italic or None)
        return par

    def p(self, text='', size=None, bold=False, color=None, align=None, after=None, before=None,
          italic=False, container=None, keep=False):
        cont = container if container is not None else self.d
        par = cont.add_paragraph()
        set_par_bidi(par._p)
        self._runs(par, text, size=size, bold=bold, color=color, italic=italic)
        if align == 'center':
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        if before is not None:
            par.paragraph_format.space_before = Pt(before)
        if keep:
            par.paragraph_format.keep_with_next = True
        return par

    def h(self, text, level=1):
        par = self.d.add_heading('', level=level)
        set_par_bidi(par._p)
        self._runs(par, text)
        return par

    def bullet(self, text, level=0, size=None, container=None):
        cont = container if container is not None else self.d
        par = cont.add_paragraph()
        set_par_bidi(par._p)
        pf = par.paragraph_format
        pf.right_indent = Cm(0.5 + 0.6 * level)
        pf.first_line_indent = Cm(-0.45)
        pf.space_after = Pt(3)
        self._runs(par, ('• ' if level == 0 else '◦ ') + text, size=size)
        return par

    def title_block(self, title, subtitle=None, kicker=None):
        if kicker:
            self.p(kicker, size=10, color='555555', align='center', after=2)
        self.p(title, size=20, bold=True, color='1F3A5F', align='center', after=4)
        if subtitle:
            self.p(subtitle, size=12, color='333333', align='center', after=10)
        self.p(MARK + ' — ' + ENTITY_FULL, size=9.5, color='A61B1B', align='center', after=12, bold=True)

    def note(self, text, fill='FFF4E5', size=10.5, bold_lead=None):
        """One-cell shaded box for notes and warnings."""
        t = self.d.add_table(rows=1, cols=1)
        t.style = self.d.styles['Table Grid']
        self._tbl_rtl(t, [17.0])
        c = t.cell(0, 0)
        shade(c, fill)
        c.paragraphs[0].text = ''
        set_par_bidi(c.paragraphs[0]._p)
        first = True
        for line in text.split('\n'):
            par = c.paragraphs[0] if first else c.add_paragraph()
            set_par_bidi(par._p)
            if first and bold_lead:
                self._runs(par, '**' + bold_lead + '** ', size=size)
            self._runs(par, line, size=size)
            par.paragraph_format.space_after = Pt(2)
            first = False
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def new_section(self, landscape=False):
        """Start a new page section (portrait or landscape); header/footer stay linked."""
        from docx.enum.section import WD_SECTION
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        _set_child(sec._sectPr, 'bidi', _SECT_ORDER)
        return sec

    def page_break(self):
        par = self.d.add_paragraph()
        par.add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ tables
    def _tbl_rtl(self, t, widths):
        tbl = t._tbl
        tblPr = tbl.tblPr
        tblPr.get_or_add_bidiVisual()
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        lay = tblPr.get_or_add_tblLayout()
        lay.set(qn('w:type'), 'fixed')
        total = sum(widths)
        tw = tblPr.find(qn('w:tblW'))
        if tw is None:
            tw = OxmlElement('w:tblW')
            tblPr.append(tw)
        tw.set(qn('w:w'), str(int(total / 2.54 * 1440)))
        tw.set(qn('w:type'), 'dxa')
        grid = tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            if i < len(widths):
                gc.set(qn('w:w'), str(int(widths[i] / 2.54 * 1440)))
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                if i < len(widths):
                    cell.width = Cm(widths[i])

    def _fill_cell(self, cell, value, size, bold=False, color=None):
        cell.paragraphs[0].text = ''
        lines = str(value if value is not None else '').split('\n')
        for i, line in enumerate(lines):
            par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            set_par_bidi(par._p)
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.05
            if line.startswith('• '):
                par.paragraph_format.right_indent = Cm(0.1)
            self._runs(par, line, size=size, bold=bold, color=color)

    def table(self, headers, rows, widths, size=9.5, head_size=None, header_fill=HEAD_FILL,
              first_col_fill=None, row_fills=None, cant_split=True, row_height=None):
        ncols = len(headers) if headers else len(rows[0])
        t = self.d.add_table(rows=(1 if headers else 0) + len(rows), cols=ncols)
        t.style = self.d.styles['Table Grid']
        self._tbl_rtl(t, widths)
        r0 = 0
        if headers:
            hdr = t.rows[0]
            trpr = hdr._tr.get_or_add_trPr()
            _set_child(trpr, 'tblHeader', _TRPR_ORDER)
            _set_child(trpr, 'cantSplit', _TRPR_ORDER)
            for i, htxt in enumerate(headers):
                c = hdr.cells[i]
                shade(c, header_fill)
                self._fill_cell(c, htxt, head_size or size, bold=True)
            r0 = 1
        for ri, row in enumerate(rows):
            tr = t.rows[r0 + ri]
            if cant_split:
                _set_child(tr._tr.get_or_add_trPr(), 'cantSplit', _TRPR_ORDER)
            if row_height:
                trh = _set_child(tr._tr.get_or_add_trPr(), 'trHeight', _TRPR_ORDER)
                trh.set(qn('w:val'), str(int(row_height / 2.54 * 1440)))
                trh.set(qn('w:hRule'), 'atLeast')
            for ci, val in enumerate(row):
                c = tr.cells[ci]
                self._fill_cell(c, val, size)
                if first_col_fill and ci == 0:
                    shade(c, first_col_fill)
                if row_fills and row_fills.get(ri):
                    shade(c, row_fills[ri])
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv(self, pairs, widths=(4.5, 12.5), size=10, label_fill=LABEL_FILL, min_height=None):
        """Two-column label/value table (labels in the right-hand column)."""
        t = self.d.add_table(rows=len(pairs), cols=2)
        t.style = self.d.styles['Table Grid']
        self._tbl_rtl(t, list(widths))
        for i, (k, v) in enumerate(pairs):
            row = t.rows[i]
            _set_child(row._tr.get_or_add_trPr(), 'cantSplit', _TRPR_ORDER)
            shade(row.cells[0], label_fill)
            self._fill_cell(row.cells[0], k, size, bold=True)
            self._fill_cell(row.cells[1], v, size)
            if min_height:
                h = min_height[i] if isinstance(min_height, (list, tuple)) else min_height
                if h:
                    trh = _set_child(row._tr.get_or_add_trPr(), 'trHeight', _TRPR_ORDER)
                    trh.set(qn('w:val'), str(int(h / 2.54 * 1440)))
                    trh.set(qn('w:hRule'), 'atLeast')
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.d.save(path)
        return path


# ---------------------------------------------------------------- xlsx
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

_THIN = Side(style='thin', color='9AA5B1')
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = 'فريق التدريب — بيانات افتراضية'
    wb.properties.title = MARK
    return wb


def add_sheet(wb, name, title, headers, rows, widths, note=None, formats=None, wrap_cols=None):
    """RTL sheet: row1 title, row2 training-data mark, row3 note, row4 header (frozen)."""
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.rightToLeft = True
    n = len(headers)
    last = get_column_letter(n)
    ws.merge_cells(f'A1:{last}1')
    ws['A1'] = ltr(title)
    ws['A1'].font = Font(name=FONT, size=14, bold=True, color='1F3A5F')
    ws['A1'].alignment = Alignment(horizontal='right', vertical='center', readingOrder=2)
    ws.row_dimensions[1].height = 24
    ws.merge_cells(f'A2:{last}2')
    ws['A2'] = FOOTER
    ws['A2'].font = Font(name=FONT, size=10, bold=True, color='A61B1B')
    ws['A2'].alignment = Alignment(horizontal='right', readingOrder=2)
    ws.merge_cells(f'A3:{last}3')
    ws['A3'] = ltr(note or '')
    ws['A3'].font = Font(name=FONT, size=9.5, italic=True, color='444444')
    ws['A3'].alignment = Alignment(horizontal='right', wrap_text=True, vertical='top', readingOrder=2)
    if note:
        ws.row_dimensions[3].height = max(18, 15 * (1 + len(note) // 140))
    hdr_fill = PatternFill('solid', fgColor=HEAD_FILL)
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=ci, value=ltr(h))
        c.font = Font(name=FONT, size=10, bold=True)
        c.fill = hdr_fill
        c.border = BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
    ws.row_dimensions[4].height = 32
    wrap_cols = set(wrap_cols or [])
    for ri, row in enumerate(rows, start=5):
        for ci, v in enumerate(row, start=1):
            c = ws.cell(row=ri, column=ci, value=ltr(v) if isinstance(v, str) else v)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = Alignment(horizontal='right', vertical='top', readingOrder=2,
                                    wrap_text=(ci - 1) in wrap_cols)
            if formats and (ci - 1) in formats:
                c.number_format = formats[ci - 1]
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = 'A5'
    if rows:
        ws.auto_filter.ref = f'A4:{last}{4 + len(rows)}'
    ws.print_title_rows = '1:4'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = FOOTER + ' — &P/&N'
    ws.oddFooter.center.size = 8
    ws.oddHeader.right.text = ltr(title)
    ws.oddHeader.right.size = 8
    return ws


# ---------------------------------------------------------------- official text
# Arabic wording written faithfully from the official Arabic + English texts
# (src_official/ecc-ar-norm.txt has reversed codes and dropped glyphs).
CTRL = {
    '1-9-5': 'يجب مراجعة وإلغاء الصلاحيات للعاملين مباشرة بعد انتهاء/إنهاء العلاقة الوظيفية مع الجهة.',
    '1-9-4-1': 'التوعية بالأمن السيبراني (عند بداية المهنة الوظيفية وخلالها).',
    '1-10-1': 'يجب تطوير واعتماد برنامج للتوعية بالأمن السيبراني في الجهة من خلال قنوات متعددة دورياً، وذلك لتعزيز الوعي بالأمن السيبراني وتهديداته ومخاطره، وبناء ثقافة إيجابية للأمن السيبراني.',
    '1-10-2': 'يجب تطبيق البرنامج المعتمد للتوعية بالأمن السيبراني في الجهة.',
    '1-10-3-1': 'التعامل الآمن مع خدمات البريد الإلكتروني خصوصاً مع رسائل التصيد الإلكتروني.',
    '1-10-5': 'يجب مراجعة تطبيق برنامج التوعية بالأمن السيبراني في الجهة دورياً.',
    '2-2-1': 'يجب تحديد وتوثيق واعتماد متطلبات الأمن السيبراني لإدارة هويات الدخول والصلاحيات في الجهة.',
    '2-2-2': 'يجب تطبيق متطلبات الأمن السيبراني لإدارة هويات الدخول والصلاحيات في الجهة.',
    '2-2-3-2': 'التحقق من الهوية متعدد العناصر (Multi-Factor Authentication) وتحديد عناصر التحقق المناسبة وعددها وكذلك تقنيات التحقق المناسبة بناءً على نتائج تقييم الأثر المحتمل لفشل عملية التحقق وتخطيها، وذلك لعمليات الدخول عن بعد والحسابات ذات الصلاحيات الهامة والحساسة.',
    '2-2-3-3': 'إدارة تصاريح وصلاحيات المستخدم (Authorization) بناءً على مبادئ التحكم بالدخول والصلاحيات (مبدأ الحاجة إلى المعرفة والاستخدام، ومبدأ الحد الأدنى من الصلاحيات والامتيازات، ومبدأ فصل المهام).',
    '2-2-3-4': 'إدارة الصلاحيات الهامة والحساسة (Privileged Access Management).',
    '2-2-3-5': 'المراجعة الدورية لهويات الدخول والصلاحيات.',
    '2-2-4': 'يجب مراجعة تطبيق متطلبات الأمن السيبراني لإدارة هويات الدخول والصلاحيات في الجهة دورياً.',
    '2-5-3-1': 'العزل والتقسيم المادي أو المنطقي لأجزاء الشبكات بشكل آمن، واللازم للسيطرة على مخاطر الأمن السيبراني ذات العلاقة، باستخدام جدار الحماية (Firewall) ومبدأ الدفاع الأمني متعدد المراحل (Defense-in-Depth).',
    '2-5-3-5': 'قيود وإدارة منافذ وبروتوكولات وخدمات الشبكة.',
    '2-5-4': 'يجب مراجعة تطبيق متطلبات الأمن السيبراني لإدارة أمن شبكات الجهة دورياً.',
    '2-9-1': 'يجب تحديد وتوثيق واعتماد متطلبات الأمن السيبراني لإدارة النسخ الاحتياطية للجهة.',
    '2-9-2': 'يجب تطبيق متطلبات الأمن السيبراني لإدارة النسخ الاحتياطية للجهة.',
    '2-9-3-1': 'نطاق النسخ الاحتياطية وشموليتها للأصول المعلوماتية والتقنية الحساسة.',
    '2-9-3-2': 'القدرة السريعة على استعادة البيانات والأنظمة بعد التعرض لحوادث الأمن السيبراني.',
    '2-9-3-3': 'إجراء فحص دوري لمدى فعالية استعادة النسخ الاحتياطية.',
    '2-9-4': 'يجب مراجعة تطبيق متطلبات الأمن السيبراني لإدارة النسخ الاحتياطية للجهة.',
    '2-12-3-2': 'تفعيل سجلات الأحداث الخاصة بالحسابات ذات الصلاحيات الهامة والحساسة على الأصول المعلوماتية وأحداث عمليات الدخول عن بعد لدى الجهة.',
    '4-1-1': 'يجب تحديد وتوثيق واعتماد متطلبات الأمن السيبراني ضمن العقود والاتفاقيات مع الأطراف الخارجية للجهة.',
    '4-1-2': 'يجب أن تغطي متطلبات الأمن السيبراني ضمن العقود والاتفاقيات (مثل اتفاقية مستوى الخدمة SLA) مع الأطراف الخارجية التي قد تتأثر بإصابتها بيانات الجهة أو الخدمات المقدمة لها بحد أدنى ما يلي:',
    '4-1-2-1': 'بنود المحافظة على سرية المعلومات (Non-Disclosure Clauses) والحذف الآمن من قِبَل الطرف الخارجي لبيانات الجهة عند انتهاء الخدمة.',
    '4-1-2-2': 'إجراءات التواصل في حال حدوث حادثة أمن سيبراني.',
    '4-1-2-3': 'إلزام الطرف الخارجي بتطبيق متطلبات وسياسات الأمن السيبراني للجهة والمتطلبات التشريعية والتنظيمية ذات العلاقة.',
    '4-1-3': 'يجب أن تغطي متطلبات الأمن السيبراني مع الأطراف الخارجية التي تقدم خدمات الإسناد للأمن السيبراني أو لتقنية المعلومات، أو الخدمات المدارة بحد أدنى ما يلي:',
    '4-1-3-1': 'إجراء تقييم لمخاطر الأمن السيبراني، والتأكد من وجود ما يضمن السيطرة على تلك المخاطر، قبل توقيع العقود والاتفاقيات أو عند تغيير المتطلبات التشريعية والتنظيمية ذات العلاقة.',
    '4-1-3-2': 'أن تكون مراكز عمليات خدمات الأمن السيبراني المدارة للتشغيل والمراقبة، والتي تستخدم طريقة الوصول عن بعد، موجودة بالكامل داخل المملكة.',
    '4-1-4': 'يجب مراجعة متطلبات الأمن السيبراني مع الأطراف الخارجية دورياً.',
    '4-2-3-1': 'حماية بيانات الجهة من قبل مقدمي خدمات الحوسبة السحابية والاستضافة بما يتوافق مع مستوى تصنيفها، وإعادتها للجهة (بصيغة قابلة للاستخدام) عند انتهاء الخدمة.',
    '4-2-3-2': 'فصل البيئة الخاصة بالجهة (وخصوصاً الخوادم الافتراضية) عن غيرها من البيئات التابعة لجهات أخرى في خدمات الحوسبة السحابية.',
}


def official(code):
    return '[[متطلب رسمي]] ' + code + ': «' + CTRL[code] + '»'


def official_codes():
    """All control codes present in the official English text (normalised)."""
    path = os.path.join(SP, 'src_official', 'ecc-en.txt')
    txt = open(path, encoding='utf-8').read()
    codes = set(re.findall(r'\b\d{1,2}-\d{1,2}(?:-\d{1,2}){0,2}\b', txt))
    codes |= {c.replace('.', '-') for c in re.findall(r'\b\d{1,2}\.\d{1,2}\.\d{1,2}\.\d{1,2}\b', txt)}
    return codes


def _count(n, one, two, few, many):
    """Arabic counted noun: 1 واحدة/2 مثنى/3-10 جمع/11+ مفرد منصوب."""
    if n == 1:
        return one
    if n == 2:
        return two
    if 3 <= n % 100 <= 10:
        return f'{n} {few}'
    return f'{n} {many}'


def fmt_minutes(m):
    m = int(round(m))
    h, mm = divmod(m, 60)
    hs = _count(h, 'ساعة واحدة', 'ساعتان', 'ساعات', 'ساعة')
    ms = _count(mm, 'دقيقة واحدة', 'دقيقتان', 'دقائق', 'دقيقة')
    if h and mm:
        return f'{hs} و{ms}'
    return hs if h else ms


def days_ar(n):
    return _count(n, 'يوم واحد', 'يومان', 'أيام', 'يوماً')


def ar_list(items):
    """Join Latin identifiers with Arabic «، و» so each stays a separate LTR run."""
    items = list(items)
    if len(items) <= 1:
        return ''.join(items)
    return items[0] + ''.join('، و' + x for x in items[1:])


def pct(a, b, nd=1):
    return f'{(100.0 * a / b):.{nd}f}%' if b else '—'
