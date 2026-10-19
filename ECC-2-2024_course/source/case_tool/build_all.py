# -*- coding: utf-8 -*-
"""Regenerate every training attachment and run automatic consistency checks.

Usage: python3 case_tool/build_all.py
Outputs: deliver/attachments/ (see README_المرفقات.docx). Visual check: python3 case_tool/render.py
"""
import glob
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import gen_00_entity, gen_01_iam, gen_02_backup, gen_03_thirdparty, gen_04_perf, gen_readme  # noqa: E402
from common import OUT, FOOTER, LRM, _RX, official_codes  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

MINE = ('00_', '01_', '02_', '03_', '04_', 'README')
TRAINER_ONLY = ('الحل_ومعيار_التصحيح', 'مفتاح_الحل', 'README')
LEAK_WORDS = ['الحل النموذجي', 'فجوة', 'فجوات', 'فخاخ', 'مقصودة', 'خطأ شائع', 'أخطاء شائعة', 'الإجابات النموذجية', 'ملاحظة جوهرية']


def docx_text(path):
    d = Document(path)
    out = []
    for p in d.paragraphs:
        out.append(p.text)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
    for s in d.sections:
        out += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return '\n'.join(out), d


def xlsx_text(path):
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out += [str(v) for v in row if isinstance(v, str)]
    return '\n'.join(out), wb


def check():
    codes = official_codes()
    problems = []
    files = sorted(glob.glob(os.path.join(OUT, '**', '*.docx'), recursive=True) +
                   glob.glob(os.path.join(OUT, '**', '*.xlsx'), recursive=True))
    files = [f for f in files if os.path.relpath(f, OUT).startswith(MINE)]  # other folders are not generated here
    used = set()
    for f in files:
        rel = os.path.relpath(f, OUT)
        if f.endswith('.docx'):
            txt, d = docx_text(f)
            for s in d.sections:
                ft = ''.join(p.text for p in s.footer.paragraphs)
                if FOOTER not in ft:
                    problems.append(f'{rel}: footer missing')
                if s.page_width.cm < 20.9 or s.page_width.cm > 29.8:
                    problems.append(f'{rel}: page size?')
        else:
            txt, wb = xlsx_text(f)
            for ws in wb.worksheets:
                if ws['A2'].value != FOOTER:
                    problems.append(f'{rel}/{ws.title}: training mark missing')
                if not ws.sheet_view.rightToLeft:
                    problems.append(f'{rel}/{ws.title}: not RTL')
                if ws.freeze_panes != 'A5':
                    problems.append(f'{rel}/{ws.title}: header not frozen')
        # every code/date wrapped with LRM on both sides
        for m in _RX.finditer(txt):
            a, b = m.start(), m.end()
            if not (a > 0 and txt[a - 1] == LRM and b < len(txt) and txt[b] == LRM):
                problems.append(f'{rel}: unwrapped "{m.group(0)}"')
        for c in re.findall(LRM + r'(\d{1,2}(?:-\d{1,2}){1,3})' + LRM, txt):
            used.add(c)
            if c not in codes:
                problems.append(f'{rel}: control code {c} not in official text')
        if not any(k in rel for k in TRAINER_ONLY):
            for w in LEAK_WORDS:
                if w in txt:
                    problems.append(f'{rel}: trainee file contains solution word «{w}»')
    return problems, used


def main():
    for g in (gen_00_entity, gen_01_iam, gen_02_backup, gen_03_thirdparty, gen_04_perf, gen_readme):
        print('built', g.build())
    problems, used = check()
    print('control codes cited:', ' '.join(sorted(used, key=lambda c: [int(x) for x in c.split('-')])))
    print('\n'.join(problems) if problems else 'checks: OK')
    for k, v in gen_04_perf.facts().items():
        print(k, v)
    return 1 if problems else 0


if __name__ == '__main__':
    sys.exit(main())
