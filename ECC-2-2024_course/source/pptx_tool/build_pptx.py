#!/usr/bin/env python3
"""HTML slide deck (deck.json + slides/<id>.html, 1920x1080, RTL Arabic) -> editable PPTX + PDF + QA.

    python3 -I build_pptx.py --deck <project dir> --blobs blob_map.json --out <out dir>/name.pptx

Pipeline (see README.md):
  1. fonts   download the Google Fonts TTFs (OFL), install to ~/.local/share/fonts, ship a copy
  2. render  Chromium (render.cjs): reference PNG per slide + DOM walk -> JSON + QA colour pass
  3. build   one native PowerPoint object per painted box / text / table / image / icon
  4. qa      LibreOffice -> PDF -> PNG, side-by-side sheets, overflow / wrap / LTR-token report
  5. the reference PDF comes straight from Chromium (step 2)
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

TOOL = Path(__file__).resolve().parent
VENDOR = TOOL / 'vendor'
EMU_PX = 6350                 # 1920 px = 12192000 EMU
PT_PX = 0.5                   # 1 px = 0.5 pt  (font size px * 0.5 = pt)
SLIDE_W, SLIDE_H = 1920, 1080
SLACK = 0.03                  # extra width on the side opposite the alignment

# ---------------------------------------------------------------- fonts (OFL, github.com/google/fonts)
GF_RAW = 'https://raw.githubusercontent.com/google/fonts/main/ofl'
FONTS = {  # family -> (repo dir, local dir, [(file, css weight)])
    'Alexandria': ('alexandria', 'Alexandria', [('Alexandria[wght].ttf', '100 900')]),
    'IBM Plex Sans Arabic': ('ibmplexsansarabic', 'IBMPlexSansArabic', [
        ('IBMPlexSansArabic-Regular.ttf', '400'), ('IBMPlexSansArabic-Medium.ttf', '500'),
        ('IBMPlexSansArabic-SemiBold.ttf', '600'), ('IBMPlexSansArabic-Bold.ttf', '700')]),
    'IBM Plex Mono': ('ibmplexmono', 'IBMPlexMono', [
        ('IBMPlexMono-Regular.ttf', '400'), ('IBMPlexMono-SemiBold.ttf', '600')]),
}
ARABIC_FONT = 'IBM Plex Sans Arabic'          # complex-script face for runs whose font has no Arabic
HAS_ARABIC = {'Alexandria', 'IBM Plex Sans Arabic'}

# ---------------------------------------------------------------- icons (Phosphor, MIT)
PHOSPHOR = 'https://unpkg.com/@phosphor-icons/core@2.1.1'
ICONS = {  # x-icon name -> Phosphor "regular" asset
    'Activity': 'pulse', 'Book': 'book', 'Chart': 'chart-bar', 'Chat': 'chat-circle', 'Check': 'check',
    'CheckCircle': 'check-circle', 'Clock': 'clock', 'Cloud': 'cloud', 'Code': 'code', 'Database': 'database',
    'Globe': 'globe', 'GraduationCap': 'graduation-cap', 'Home': 'house', 'Key': 'key', 'Lightbulb': 'lightbulb',
    'Lightning': 'lightning', 'Link': 'link', 'Lock': 'lock', 'PaperPlane': 'paper-plane-tilt', 'Play': 'play',
    'Search': 'magnifying-glass', 'Settings': 'gear', 'Star': 'star', 'ThumbsUp': 'thumbs-up', 'Tool': 'wrench',
    'Trust': 'handshake', 'Users': 'users', 'Verified': 'seal-check', 'Warning': 'warning', 'Wrench': 'wrench',
}

ARABIC_RE = re.compile('[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]')
RLE = '\u202B'                # stripped from text; every other bidi mark (LRM, ...) is kept
TOKEN_RE = re.compile(r'[A-Za-z0-9]+(?:[-:/.,%]+[A-Za-z0-9]+)*')
LRM = '\u200E'

# QA colours: channels in steps of 36 with at least one channel 0, so anti-aliased
# blends with the white background never look like another QA colour.
QA_STEP = 36
QA_COLORS = [(r, g, b) for r in range(0, 217, 36) for g in range(0, 217, 36) for b in range(0, 217, 36)
             if min(r, g, b) == 0 and (r, g, b) != (0, 0, 0)]


def log(*a):
    print(*a, flush=True)


# ================================================================= step 1: fonts + icons
def fetch(url, dest):
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_name(dest.name + '.part')
    if shutil.which('curl'):
        subprocess.run(['curl', '-sSfL', '--retry', '3', '-o', str(tmp), url], check=True)
    else:
        with urllib.request.urlopen(url, timeout=60) as r, open(tmp, 'wb') as f:
            shutil.copyfileobj(r, f)
    tmp.replace(dest)


def is_font(p):
    with open(p, 'rb') as f:
        return f.read(4) in (b'\x00\x01\x00\x00', b'true', b'OTTO')


def setup_fonts(out_dir):
    """Download (once, into vendor/fonts), install for fontconfig, copy next to the deliverable."""
    vend = VENDOR / 'fonts'
    inst = Path.home() / '.local' / 'share' / 'fonts' / 'deck-fonts'
    inst.mkdir(parents=True, exist_ok=True)
    ship = out_dir / 'fonts'
    files = {}
    for fam, (repo, local, faces) in FONTS.items():
        d = vend / local
        lic = d / 'OFL.txt'
        if not lic.exists():
            fetch(f'{GF_RAW}/{repo}/OFL.txt', lic)
        for fn, w in faces:
            p = d / fn
            if not p.exists():
                log(f'  download {fn}')
                fetch(f'{GF_RAW}/{repo}/' + fn.replace('[', '%5B').replace(']', '%5D'), p)
            if not is_font(p):
                raise SystemExit(f'not a font file: {p}')
            shutil.copy2(p, inst / fn)
            files.setdefault(fam, []).append((p, w))
            if w in ('400', '100 900'):
                FONT_K[fam] = read_metrics(p)
        (ship / local).mkdir(parents=True, exist_ok=True)
        for fn, _ in faces:
            shutil.copy2(d / fn, ship / local / fn)
        shutil.copy2(lic, ship / local / 'OFL.txt')
    subprocess.run(['fc-cache', '-f', str(inst.parent)], check=False, stdout=subprocess.DEVNULL)
    return files


def font_css(files):
    css = []
    for fam, faces in files.items():
        for p, w in faces:
            css.append(f"@font-face {{ font-family: '{fam}'; src: url('{p.as_uri()}'); font-weight: {w}; font-style: normal; }}")
    # Plex Mono has no Arabic: let its Arabic characters use Plex Sans Arabic (as the PPTX does)
    ar = dict((w, p) for p, w in files.get('IBM Plex Sans Arabic', []))
    for w in ('400', '600'):
        if w in ar:
            css.append(f"@font-face {{ font-family: 'IBM Plex Mono'; src: url('{ar[w].as_uri()}'); font-weight: {w}; "
                       "unicode-range: U+0600-06FF, U+0750-077F, U+08A0-08FF, U+FB50-FDFF, U+FE70-FEFF; }")
    return '\n'.join(css)


def setup_icons():
    d = VENDOR / 'icons'
    out = {}
    if not (d / 'LICENSE-phosphor.txt').exists():
        fetch(f'{PHOSPHOR}/LICENSE', d / 'LICENSE-phosphor.txt')
    for name, asset in ICONS.items():
        p = d / f'{asset}.svg'
        if not p.exists():
            log(f'  download icon {asset}')
            fetch(f'{PHOSPHOR}/assets/regular/{asset}.svg', p)
        head = p.read_text('utf-8', errors='replace')[:400]
        if '<svg' not in head or '<script' in p.read_text('utf-8', errors='replace'):
            raise SystemExit(f'unexpected icon file {p}')
        out[name] = str(p)
    return out


# ================================================================= step 2: Chromium render
def find_node_path():
    node = shutil.which('node')
    if not node:
        raise SystemExit('node not found (needed for Playwright)')
    env = dict(os.environ)
    try:
        root = subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True, check=True).stdout.strip()
    except Exception:
        root = str(Path(node).resolve().parent.parent / 'lib' / 'node_modules')
    env['NODE_PATH'] = os.pathsep.join(x for x in [env.get('NODE_PATH', ''), root] if x)
    if not env.get('PLAYWRIGHT_BROWSERS_PATH') and Path('/opt/pw-browsers').exists():
        env['PLAYWRIGHT_BROWSERS_PATH'] = '/opt/pw-browsers'
    return node, env


def render(slides, deck_dir, blobs, fcss, icons, work, ref_dir, ref_pdf):
    job = {
        'slides': [], 'blobs': blobs, 'fontCss': fcss, 'icons': icons,
        'qaColors': [f'rgb({r}, {g}, {b})' for r, g, b in QA_COLORS],
        'workDir': str(work), 'iconDir': str(work / 'icons'), 'outJson': str(work / 'dom.json'),
        'refPdf': str(ref_pdf) if ref_pdf else None,
    }
    (work / 'icons').mkdir(parents=True, exist_ok=True)
    (work / 'diag_ref').mkdir(parents=True, exist_ok=True)
    for i, sid in enumerate(slides, 1):
        job['slides'].append({
            'id': sid, 'index': i, 'file': str(deck_dir / 'slides' / f'{sid}.html'),
            'refPng': str(ref_dir / f'{i:03d}-{sid}.png'), 'diagPng': str(work / 'diag_ref' / f'{i:03d}-{sid}.png'),
        })
    jf = work / 'job.json'
    jf.write_text(json.dumps(job, ensure_ascii=False), 'utf-8')
    node, env = find_node_path()
    subprocess.run([node, str(TOOL / 'render.cjs'), str(jf)], check=True, env=env)
    return json.loads((work / 'dom.json').read_text('utf-8'))['slides']


# ================================================================= helpers
def parse_color(s):
    """CSS computed colour -> (r, g, b, a) or None."""
    if not s:
        return None
    s = s.strip()
    m = re.match(r'rgba?\(([^)]*)\)', s)
    if m:
        parts = [p for p in re.split(r'[,\s/]+', m.group(1).strip()) if p]
        r, g, b = (int(round(float(x))) for x in parts[:3])
        a = float(parts[3]) if len(parts) > 3 else 1.0
        return (r, g, b, a)
    m = re.match(r'#([0-9a-fA-F]{6})$', s)
    if m:
        v = m.group(1)
        return (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16), 1.0)
    return None


def emu(v):
    return Emu(int(round(v * EMU_PX)))


def first_family(ff):
    fam = (ff or '').split(',')[0].strip().strip('"\'')
    return fam or ARABIC_FONT


def set_color(parent_fill_xml, rgba, extra_alpha=1.0):
    """parent_fill_xml: an <a:solidFill>; set srgbClr + alpha."""
    for ch in list(parent_fill_xml):
        parent_fill_xml.remove(ch)
    clr = etree.SubElement(parent_fill_xml, qn('a:srgbClr'))
    clr.set('val', '%02X%02X%02X' % tuple(rgba[:3]))
    a = rgba[3] * extra_alpha
    if a < 0.999:
        al = etree.SubElement(clr, qn('a:alpha'))
        al.set('val', str(int(round(a * 100000))))


def spPr_of(shape):
    return shape._element.spPr


def fill_shape(shape, rgba, opacity=1.0):
    if rgba is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*rgba[:3])
    set_color(spPr_of(shape).find(qn('a:solidFill')), rgba, opacity)


def line_shape(shape, rgba=None, width_px=0, style='solid', opacity=1.0):
    if rgba is None or width_px <= 0:
        shape.line.fill.background()
        return
    shape.line.width = emu(width_px)
    shape.line.color.rgb = RGBColor(*rgba[:3])
    ln = spPr_of(shape).find(qn('a:ln'))
    set_color(ln.find(qn('a:solidFill')), rgba, opacity)
    if style == 'dashed':
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    elif style == 'dotted':
        shape.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT


def drop_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def corner_radii(item, w, h):
    rs = []
    for q in item.get('radius') or []:
        if q.get('pct') is not None:
            rs.append(('pct', q['pct']))
        else:
            rs.append(('px', min(q.get('px', 0), w / 2, h / 2)))
    while len(rs) < 4:
        rs.append(('px', 0))
    return rs


def intersect(a, b):
    x0, y0 = max(a[0], b[0]), max(a[1], b[1])
    x1, y1 = min(a[0] + a[2], b[0] + b[2]), min(a[1] + a[3], b[1] + b[3])
    return (x0, y0, max(0, x1 - x0), max(0, y1 - y0))


# ================================================================= text runs
def collapse_runs(runs):
    """CSS white-space:normal collapsing across runs; strips bidi embedding controls."""
    out, prev_space = [], True
    for r in runs:
        if r.get('br'):
            out.append({'br': True})
            prev_space = True
            continue
        had_ctrl = RLE in r['t']
        t = r['t'].replace(RLE, '')
        if r.get('ws') not in ('pre', 'pre-wrap', 'break-spaces'):
            t = re.sub(r'[ \t\n\r\f]+', ' ', t)
            if prev_space and t.startswith(' '):
                t = t[1:]
        if r.get('tt') == 'uppercase':
            t = t.upper()
        elif r.get('tt') == 'lowercase':
            t = t.lower()
        if not t:
            if had_ctrl:
                out.append({**r, 't': '', 'rle': True})
            continue
        prev_space = t.endswith(' ')
        out.append({**r, 't': t, 'rle': had_ctrl})
    while out and not out[-1].get('br') and out[-1]['t'].endswith(' '):
        out[-1]['t'] = out[-1]['t'].rstrip(' ')
        if not out[-1]['t']:
            out.pop()
    return [r for r in out if r.get('br') or r['t']]


def is_rtl(raw_runs, runs):
    if any(RLE in r.get('t', '') for r in raw_runs):
        return True
    return any(ARABIC_RE.search(r.get('t', '')) for r in runs if not r.get('br'))


def ltr_tokens(text):
    return [m.group(0) for m in TOKEN_RE.finditer(text) if re.search(r'[-:/]', m.group(0))]


def protect_tokens(runs):
    """Insert LRM before/after every Latin/numeric token holding - : or / (spans runs).
    Without it, digit codes after Arabic letters become Arabic numbers (UBA W2) and 2-9-1 shows as 1-9-2."""
    full = ''.join(r.get('t', '') for r in runs)
    ins = {}
    for m in TOKEN_RE.finditer(full):
        if re.search(r'[-:/]', m.group(0)):
            if m.start() == 0 or full[m.start() - 1] != LRM:       # the source may already carry LRMs
                ins.setdefault(m.start(), []).append('s')
            if m.end() >= len(full) or full[m.end()] != LRM:
                ins.setdefault(m.end() - 1, []).append('e')
    if not ins:
        return runs
    out, pos = [], 0
    for r in runs:
        if r.get('br'):
            out.append(r)
            continue
        t, buf = r['t'], []
        for i, ch in enumerate(t):
            marks = ins.get(pos + i, [])
            if 's' in marks:
                buf.append(LRM)
            buf.append(ch)
            if 'e' in marks:
                buf.append(LRM)
        pos += len(t)
        out.append({**r, 't': ''.join(buf)})
    return out


def pptx_align(css_align, rtl):
    a = (css_align or '').lower()
    if a in ('start', '-webkit-start'):
        a = 'right' if rtl else 'left'
    elif a in ('end', '-webkit-end'):
        a = 'left' if rtl else 'right'
    return {'right': PP_ALIGN.RIGHT, 'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
            '-webkit-center': PP_ALIGN.CENTER, 'justify': PP_ALIGN.JUSTIFY}.get(a, PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT), a


def host_span(it):
    """Horizontal extent the text frame may grow into: the painted ancestor (card) or the slide."""
    h = it.get('host') or {'x': 0, 'w': SLIDE_W}
    pad = 4.0 if (h['w'] < SLIDE_W - 1) else 0.0
    return min(it['x'], h['x'] + pad), max(it['x'] + it['w'], h['x'] + h['w'] - pad)


def slack_px(it, content_w):
    """Extra width granted on the side opposite the alignment. Single-line texts get more, so a
    label that renders a little wider (600 weight drawn with the 700 face) never breaks mid-word."""
    fs = it.get('fs', 24)
    if (it.get('nlines') or 1) <= 1:
        return max(0.08 * content_w, 0.6 * fs, 2.0)
    return max(SLACK * content_w, 0.25 * fs, 2.0)


def add_run(p, r, qa_rgb=None):
    if r.get('br'):
        p.add_line_break()
        return
    run = p.add_run()
    run.text = r['t']
    f = run.font
    f.size = Pt(round(r['size'] * PT_PX * 2) / 2)
    f.bold = r['weight'] >= 600
    f.italic = bool(r.get('italic'))
    if r.get('ul'):
        f.underline = True
    col = parse_color(r.get('color')) or (0, 0, 0, 1)
    if qa_rgb is not None:
        col = (*qa_rgb, 1.0)
    f.color.rgb = RGBColor(*col[:3])
    rPr = run._r.get_or_add_rPr()
    set_color(rPr.find(qn('a:solidFill')), col)
    if r.get('strike'):
        rPr.set('strike', 'sngStrike')
    ls = r.get('ls') or 'normal'
    if ls != 'normal':
        try:
            v = float(ls.replace('px', ''))
            if abs(v) > 0.01:
                rPr.set('spc', str(int(round(v * PT_PX * 100))))
        except ValueError:
            pass
    fam = first_family(r.get('family'))
    f.name = fam
    cs_face = fam if (fam in HAS_ARABIC or not ARABIC_RE.search(r['t'])) else ARABIC_FONT
    latin = rPr.find(qn('a:latin'))
    ea = etree.Element(qn('a:ea'))
    ea.set('typeface', fam)
    cs = etree.Element(qn('a:cs'))
    cs.set('typeface', cs_face)
    latin.addnext(ea)
    ea.addnext(cs)


def fill_text_frame(tf, it, qa_rgb=None):
    """One paragraph (the element) with its runs."""
    runs = collapse_runs(it['runs'])
    rtl = is_rtl(it['runs'], runs)
    if rtl:
        runs = protect_tokens(runs)
    tf.clear()
    p = tf.paragraphs[0]
    align, _ = pptx_align(it.get('align'), rtl)
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    if rtl:
        pPr.set('rtl', '1')
    lh = it.get('lh') or it.get('fs', 24) * 1.4
    p.line_spacing = Pt(round(lh * PT_PX, 2))
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    for r in runs:
        add_run(p, r, qa_rgb)
    if not runs:   # keep an editable empty paragraph with the right size
        end = p._p.get_or_add_endParaRPr()
        end.set('sz', str(int(round(it.get('fs', 24) * PT_PX * 100))))
    return rtl, align


def setup_frame(tf, insets, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    t, r, b, l = (max(0.0, v) for v in insets)
    tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left = emu(t), emu(r), emu(b), emu(l)
    tf.vertical_anchor = anchor
    bp = tf._txBody.find(qn('a:bodyPr'))
    bp.set('rtlCol', '0')
    bp.set('anchorCtr', '0')
    for tag in ('a:spAutoFit', 'a:normAutofit'):
        e = bp.find(qn(tag))
        if e is not None:
            bp.remove(e)
    if bp.find(qn('a:noAutofit')) is None:
        etree.SubElement(bp, qn('a:noAutofit'))


# ================================================================= step 3: PPTX objects
class SlideBuilder:
    def __init__(self, prs, slide, data, diag=False):
        self.prs, self.slide, self.data, self.diag = prs, slide, data, diag
        self.shapes = slide.shapes
        self.n = 0

    def name(self, kind):
        self.n += 1
        return f'{kind} {self.n}'

    # ------------------------------------------------------------ boxes
    def add_rect(self, x, y, w, h, rgba, opacity=1.0, radii=None, line=None, name='Box'):
        """rectangle / rounded rectangle / oval / top- or bottom-rounded rectangle."""
        if w <= 0.05 or h <= 0.05:
            return None
        geom, adj, flip = MSO_SHAPE.RECTANGLE, None, False
        if radii:
            vals = [(v if k == 'px' else min(w, h) * v / 100) for k, v in radii]
            pct = any(k == 'pct' and v >= 50 for k, v in radii)
            m = min(w, h)
            if max(vals) > 0.5:
                if pct and abs(w - h) > 1 and all(k == 'pct' for k, _ in radii):
                    geom = MSO_SHAPE.OVAL
                elif max(vals) - min(vals) <= 1:
                    geom, adj = MSO_SHAPE.ROUNDED_RECTANGLE, [min(0.5, vals[0] / m)]
                elif vals[2] <= 0.5 and vals[3] <= 0.5 and abs(vals[0] - vals[1]) <= 1:
                    geom, adj = MSO_SHAPE.ROUND_2_SAME_RECTANGLE, [min(0.5, vals[0] / m), 0.0]
                elif vals[0] <= 0.5 and vals[1] <= 0.5 and abs(vals[2] - vals[3]) <= 1:
                    geom, adj, flip = MSO_SHAPE.ROUND_2_SAME_RECTANGLE, [min(0.5, vals[2] / m), 0.0], True
                else:
                    geom, adj = MSO_SHAPE.ROUNDED_RECTANGLE, [min(0.5, sum(vals) / 4 / m)]
        lw = line[1] if line else 0
        sx, sy, sw, sh = x + lw / 2, y + lw / 2, w - lw, h - lw      # CSS border sits inside the box
        if sw <= 0 or sh <= 0:
            sx, sy, sw, sh = x, y, w, h
        # the preset's own text rectangle (presetShapeDefinitions.xml): text is laid out inside it
        ss = min(sw, sh)
        if geom == MSO_SHAPE.ROUNDED_RECTANGLE:
            il = ss * adj[0] * 0.29289
            self.text_rect = (il, il, il, il)
        elif geom == MSO_SHAPE.ROUND_2_SAME_RECTANGLE:
            il = ss * adj[0] * 0.29289
            self.text_rect = (il, 0, il, il) if flip else (il, il, il, 0)
        elif geom == MSO_SHAPE.OVAL:
            self.text_rect = (sw * 0.14645, sh * 0.14645, sw * 0.14645, sh * 0.14645)
        else:
            self.text_rect = (0, 0, 0, 0)
        shp = self.shapes.add_shape(geom, emu(sx), emu(sy), emu(sw), emu(sh))
        drop_style(shp)
        shp.name = self.name(name)
        if adj:
            for i, v in enumerate(adj):
                shp.adjustments[i] = v
        if flip:
            shp._element.spPr.find(qn('a:xfrm')).set('flipV', '1')
        fill_shape(shp, rgba, opacity)
        if line:
            line_shape(shp, line[0], line[1], line[2], opacity)
        else:
            line_shape(shp, None)
        return shp

    def box_geometry(self, it):
        x, y, w, h = it['x'], it['y'], it['w'], it['h']
        radii = corner_radii(it, w, h)
        clip = it.get('clip')
        if clip:
            cx, cy, cw, ch = intersect((x, y, w, h), (clip['x'], clip['y'], clip['w'], clip['h']))
            if cw <= 0 or ch <= 0:
                return None
            # corners that coincide with the clipping box take its rounding
            cr = clip.get('radius') or [0, 0, 0, 0]
            corners = [(cx, cy), (cx + cw, cy), (cx + cw, cy + ch), (cx, cy + ch)]
            ccorn = [(clip['x'], clip['y']), (clip['x'] + clip['w'], clip['y']),
                     (clip['x'] + clip['w'], clip['y'] + clip['h']), (clip['x'], clip['y'] + clip['h'])]
            for i in range(4):
                if abs(corners[i][0] - ccorn[i][0]) < 1.5 and abs(corners[i][1] - ccorn[i][1]) < 1.5:
                    own = radii[i][1] if radii[i][0] == 'px' else min(w, h) * radii[i][1] / 100
                    if cr[i] > own:
                        radii[i] = ('px', min(cr[i], cw / 2, ch / 2))
            x, y, w, h = cx, cy, cw, ch
        return x, y, w, h, radii

    def emit_box(self, it, with_fill=True, host=False):
        """Painted div: fill + uniform outline as one shape, single-side borders as thin rules.
        host=True always returns a shape (to carry text), even when nothing is painted."""
        g = self.box_geometry(it)
        if not g:
            return None, 0, (0, 0, 0, 0)
        x, y, w, h, radii = g
        op = it.get('opacity', 1.0)
        fill = parse_color(it.get('fill')) if with_fill else None
        b = it.get('borders') or {}
        sides = ['top', 'right', 'bottom', 'left']
        uniform = (len(b) == 4 and len({(b[s]['w'], b[s]['color'], b[s]['style']) for s in sides}) == 1)
        line = None
        if uniform and with_fill:
            line = (parse_color(b['top']['color']), b['top']['w'], b['top']['style'])
        shp = None
        self.text_rect = (0, 0, 0, 0)
        if fill or line or host:
            shp = self.add_rect(x, y, w, h, fill, op, radii, line, name='Text box' if host else 'Box')
        trect = self.text_rect
        if not uniform and with_fill:
            for s in sides:
                if s not in b:
                    continue
                bw, col = b[s]['w'], parse_color(b[s]['color'])
                if not col or col[3] <= 0:
                    continue
                r = {'top': (x, y, w, bw), 'bottom': (x, y + h - bw, w, bw),
                     'left': (x, y, bw, h), 'right': (x + w - bw, y, bw, h)}[s]
                self.add_rect(*r, col, op, None, None, name='Rule')
        return shp, (line[1] if line else 0), trect

    # ------------------------------------------------------------ text
    def emit_text(self, it):
        x, y, w, h = it['x'], it['y'], it['w'], it['h']
        b = it.get('borders') or {}
        bw = {s: (b[s]['w'] if s in b else 0) for s in ('top', 'right', 'bottom', 'left')}
        pad = it.get('pad') or [0, 0, 0, 0]
        ins = [pad[0] + bw['top'], pad[1] + bw['right'], pad[2] + bw['bottom'], pad[3] + bw['left']]
        runs = collapse_runs(it['runs'])
        painted = bool(it.get('fill')) or bool(b)
        if not runs and not painted:
            return
        if not runs:
            if not self.diag:
                self.emit_box(it)
            return
        rtl = is_rtl(it['runs'], runs)
        _, a = pptx_align(it.get('align'), rtl)
        cw = max(1.0, w - ins[1] - ins[3])
        slack = slack_px(it, cw)
        if a == 'right' or (a in ('start', 'justify', '') and rtl):
            side = 'left'
        elif a in ('center', '-webkit-center'):
            side = 'both'
        else:
            side = 'right'
        dy = CAL_DY(it)
        anchor = MSO_ANCHOR.TOP
        if painted:
            # a pill / badge / ruled paragraph: the text lives inside its own shape
            content_h = h - ins[0] - ins[2]
            shp, lw, tr = self.emit_box(it, with_fill=not self.diag, host=True)
            if shp is None:
                return
            # CSS padding is measured from the box; the preset lays text out in its own text rect
            ins = [ins[0] - lw / 2 - tr[1], ins[1] - lw / 2 - tr[2], ins[2] - lw / 2 - tr[3], ins[3] - lw / 2 - tr[0]]
            # widen the text area on the side opposite the alignment by eating padding
            if side in ('left', 'both'):
                ins[3] -= min(slack if side == 'left' else slack / 2, max(0, ins[3]))
            if side in ('right', 'both'):
                ins[1] -= min(slack if side == 'right' else slack / 2, max(0, ins[1]))
            ins[0] += dy
            ins[2] -= dy
            lh = it.get('lh') or it.get('fs', 24) * 1.4
            if abs(content_h - max(1, it.get('nlines') or 1) * lh) < 0.5 * lh:
                # text fills the content box: centre it (keeps it centred even if insets clamp)
                anchor = MSO_ANCHOR.MIDDLE
                m = min(ins[0], ins[2])
                if m < 0:
                    ins[0] -= m
                    ins[2] -= m
            frame = shp
        else:
            fx, fw = x, w
            hx0, hx1 = host_span(it)
            if side in ('left', 'both'):
                s = slack if side == 'left' else slack / 2
                take = min(s, ins[3]); ins[3] -= take; s -= take
                s = max(0.0, min(s, fx - hx0))          # never beyond the card / slide edge
                fx -= s; fw += s
            if side in ('right', 'both'):
                s = slack if side == 'right' else slack / 2
                take = min(s, ins[1]); ins[1] -= take; s -= take
                s = max(0.0, min(s, hx1 - (fx + fw)))
                fw += s
            frame = self.shapes.add_textbox(emu(fx), emu(y + dy), emu(fw), emu(h))
            frame.name = self.name('Text')
        setup_frame(frame.text_frame, ins, anchor)
        qa = parse_color(it.get('qa')) if (self.diag and it.get('qa')) else None
        fill_text_frame(frame.text_frame, it, qa[:3] if qa else (None if not self.diag else (255, 255, 255)))

    # ------------------------------------------------------------ tables
    def emit_table(self, it):
        rows = [r for r in it['rows'] if r['cells']]
        if not rows:
            return
        # visual geometry: columns from the x-positions where cells render
        xs = sorted({round(c['x'], 1) for r in rows for c in r['cells']})
        cols = []
        for v in xs:
            if not cols or v - cols[-1] > 2:
                cols.append(v)
        right = max(c['x'] + c['w'] for r in rows for c in r['cells'])
        edges = cols + [right]
        widths = [edges[i + 1] - edges[i] for i in range(len(cols))]
        x0, y0 = edges[0], rows[0]['y']
        heights = [r['h'] for r in rows]
        gs = self.shapes.add_table(len(rows), len(cols), emu(x0), emu(y0), emu(sum(widths)), emu(sum(heights)))
        gs.name = self.name('Table')
        tbl = gs.table
        tblPr = tbl._tbl.tblPr
        for k in ('firstRow', 'bandRow', 'lastRow', 'firstCol', 'lastCol', 'bandCol'):
            tblPr.set(k, '0')
        sid = tblPr.find(qn('a:tableStyleId'))
        if sid is None:
            sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
        sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'   # "No Style, No Grid"
        for i, wv in enumerate(widths):
            tbl.columns[i].width = emu(wv)
        for i, hv in enumerate(heights):
            tbl.rows[i].height = emu(hv)
        op = it.get('opacity', 1.0)
        for ri, r in enumerate(rows):
            cells = sorted(r['cells'], key=lambda c: c['x'])
            for c in cells:
                ci = min(range(len(cols)), key=lambda k: abs(cols[k] - c['x']))
                cell = tbl.cell(ri, ci)
                self.fill_cell(cell, c, r, op)

    def fill_cell(self, cell, c, row, op):
        tf = cell.text_frame
        pad = c.get('pad') or [0, 0, 0, 0]
        ins = list(pad)
        runs = collapse_runs(c['runs'])
        rtl = is_rtl(c['runs'], runs)
        _, a = pptx_align(c.get('align'), rtl)
        cw = max(1.0, c['w'] - pad[1] - pad[3])
        slack = slack_px(c, cw)
        if a == 'right' or (a in ('start', 'justify') and rtl):
            ins[3] = max(0.0, ins[3] - slack)
        elif a == 'center':
            ins[1] = max(0.0, ins[1] - slack / 2); ins[3] = max(0.0, ins[3] - slack / 2)
        else:
            ins[1] = max(0.0, ins[1] - slack)
        dy = CAL_DY(c, cell=True)
        ins[0] += dy
        ins[2] -= dy
        m = min(ins[0], ins[2])
        if m < 0:
            ins[0] -= m
            ins[2] -= m
        cell.margin_top, cell.margin_right, cell.margin_bottom, cell.margin_left = (emu(v) for v in ins)
        va = c.get('valign') or 'middle'
        cell.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'bottom': MSO_ANCHOR.BOTTOM}.get(va, MSO_ANCHOR.MIDDLE)
        tf.word_wrap = True
        qa = parse_color(c.get('qa')) if (self.diag and c.get('qa')) else None
        fill_text_frame(tf, c, qa[:3] if qa else (None if not self.diag else (255, 255, 255)))
        # borders then fill, in tcPr schema order
        tcPr = cell._tc.get_or_add_tcPr()
        for ch in list(tcPr):
            tcPr.remove(ch)
        b = c.get('borders') or {}
        for side, tag in (('left', 'a:lnL'), ('right', 'a:lnR'), ('top', 'a:lnT'), ('bottom', 'a:lnB')):
            ln = etree.SubElement(tcPr, qn(tag))
            col = parse_color(b[side]['color']) if side in b else None
            if col and not self.diag and col[3] > 0:
                ln.set('w', str(int(round(b[side]['w'] * EMU_PX))))
                set_color(etree.SubElement(ln, qn('a:solidFill')), col, op)
            else:
                ln.set('w', '0')
                etree.SubElement(ln, qn('a:noFill'))
        bg = parse_color(c.get('fill')) or parse_color(row.get('bg'))
        if bg and not self.diag:
            set_color(etree.SubElement(tcPr, qn('a:solidFill')), bg, op)
        else:
            etree.SubElement(tcPr, qn('a:noFill'))

    # ------------------------------------------------------------ pictures
    def emit_img(self, it, blobs):
        path = blobs.get(it.get('blob') or '')
        if not path or not os.path.exists(path):
            log(f'    ! missing image {it.get("blob")}')
            return
        x, y, w, h = it['x'], it['y'], it['w'], it['h']
        nw, nh = it.get('nw') or 0, it.get('nh') or 0
        if not nw or not nh:
            with Image.open(path) as im:
                nw, nh = im.size
        fit = it.get('fit') or 'fill'
        px_, py_ = 0.5, 0.5
        m = re.match(r'([\d.]+)%\s+([\d.]+)%', it.get('pos') or '')
        if m:
            px_, py_ = float(m.group(1)) / 100, float(m.group(2)) / 100
        path = ensure_supported_image(path)
        crop = None
        if fit in ('contain', 'scale-down'):
            s = min(w / nw, h / nh)
            if fit == 'scale-down':
                s = min(s, 1.0)
            dw, dh = nw * s, nh * s
            x, y, w, h = x + (w - dw) * px_, y + (h - dh) * py_, dw, dh
        elif fit == 'cover':
            s = max(w / nw, h / nh)
            dw, dh = nw * s, nh * s
            ox, oy = (dw - w) * px_, (dh - h) * py_
            crop = (ox / dw, (dw - w - ox) / dw, oy / dh, (dh - h - oy) / dh)
        pic = self.shapes.add_picture(path, emu(x), emu(y), emu(w), emu(h))
        pic.name = self.name('Picture')
        if crop:
            pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = crop
        radii = corner_radii(it, it['w'], it['h'])
        r = max(v for k, v in radii if k == 'px') if any(k == 'px' for k, _ in radii) else 0
        fills_box = abs(w - it['w']) < 1.5 and abs(h - it['h']) < 1.5
        if r > 0.5 and fills_box:
            geom = pic._element.spPr.find(qn('a:prstGeom'))
            geom.set('prst', 'roundRect')
            av = geom.find(qn('a:avLst'))
            if av is None:
                av = etree.SubElement(geom, qn('a:avLst'))
            gd = etree.SubElement(av, qn('a:gd'))
            gd.set('name', 'adj')
            gd.set('fmla', 'val %d' % int(round(min(0.5, r / min(w, h)) * 100000)))
        self.pic_opacity(pic, it.get('opacity', 1.0))

    def pic_opacity(self, pic, op):
        if op < 0.999:
            blip = pic._element.find('.//' + qn('a:blip'))
            am = etree.Element(qn('a:alphaModFix'))
            am.set('amt', str(int(round(op * 100000))))
            blip.insert(0, am)          # before any a:extLst

    def emit_icon(self, it):
        png = it.get('png')
        if not png:
            return
        pic = self.shapes.add_picture(png, emu(it['x']), emu(it['y']), emu(it['w']), emu(it['h']))
        pic.name = self.name(f'Icon {it.get("name")}')
        self.pic_opacity(pic, it.get('opacity', 1.0))

    def emit_xshape(self, it):
        x, y, w, h = it['x'], it['y'], it['w'], it['h']
        kind = it.get('kind')
        col = parse_color(it.get('color'))
        ss = min(w, h)
        if kind in ('arrow-left', 'arrow-right'):
            geom = MSO_SHAPE.LEFT_ARROW if kind == 'arrow-left' else MSO_SHAPE.RIGHT_ARROW
            adj = [0.5, min(w / ss, 0.4 * w / ss)]
        elif kind in ('arrow-up', 'arrow-down'):
            geom = MSO_SHAPE.UP_ARROW if kind == 'arrow-up' else MSO_SHAPE.DOWN_ARROW
            adj = [0.5, min(h / ss, 0.4 * h / ss)]
        elif kind == 'ellipse':
            geom, adj = MSO_SHAPE.OVAL, []
        elif kind == 'diamond':
            geom, adj = MSO_SHAPE.DIAMOND, []
        else:
            geom, adj = MSO_SHAPE.RECTANGLE, []
        shp = self.shapes.add_shape(geom, emu(x), emu(y), emu(w), emu(h))
        drop_style(shp)
        shp.name = self.name(f'Shape {kind}')
        for i, v in enumerate(adj):
            shp.adjustments[i] = v
        fill_shape(shp, col, it.get('opacity', 1.0))
        line_shape(shp, None)

    # ------------------------------------------------------------ slide
    def build(self, blobs):
        d = self.data
        bg = parse_color(d.get('bg')) or (255, 255, 255, 1)
        if self.diag:
            bg = (255, 255, 255, 1)
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = RGBColor(*bg[:3])
        for it in d['items']:
            t = it['type']
            if t == 'text':
                self.emit_text(it)
            elif t == 'table':
                self.emit_table(it)
            elif self.diag:
                continue
            elif t == 'box':
                self.emit_box(it)
            elif t == 'img':
                self.emit_img(it, blobs)
            elif t == 'icon':
                self.emit_icon(it)
            elif t == 'shape':
                self.emit_xshape(it)
        if not self.diag and d.get('hidden'):
            self.slide._element.set('show', '0')
        # data-transition = how a slide LEAVES; PowerPoint transitions belong to the slide that enters
        enter = d.get('enter')
        if not self.diag and enter in ('fade', 'push', 'magic'):
            tr = etree.SubElement(self.slide._element, qn('p:transition'))
            tr.set('spd', 'med')
            etree.SubElement(tr, qn('p:push') if enter == 'push' else qn('p:fade'))
            clr = self.slide._element.find(qn('p:clrMapOvr'))
            if clr is not None:
                clr.addnext(tr)
        if not self.diag:
            notes = re.sub(r'\s+', ' ', (d.get('notes') or '').replace(RLE, '')).strip()
            if ARABIC_RE.search(notes):
                notes = protect_tokens([{'t': notes}])[0]['t']
            if notes:
                ntf = self.slide.notes_slide.notes_text_frame
                ntf.text = notes
                for p in ntf.paragraphs:
                    if ARABIC_RE.search(notes):
                        p._p.get_or_add_pPr().set('rtl', '1')
                        p.alignment = PP_ALIGN.RIGHT
                    for r in p.runs:
                        r.font.name = ARABIC_FONT
                        rPr = r._r.get_or_add_rPr()
                        cs = etree.SubElement(rPr, qn('a:cs'))
                        cs.set('typeface', ARABIC_FONT)


# ---------------------------------------------------------------- vertical calibration
# CSS puts the baseline of a line box of height L at  L/2 + (ascent - descent)/2 * fs  (half-leading
# model, hhea/typo metrics). LibreOffice with exact line spacing (spcPts) puts it at  L - LO_DESCENT * fs
# whatever the font (measured, see README). Text frames are moved by the difference.
LO_DESCENT = 0.21
FONT_K = {}          # family -> (ascent - descent) / 2 / unitsPerEm, read from the TTFs


def read_metrics(path):
    import struct
    b = Path(path).read_bytes()
    n = struct.unpack('>H', b[4:6])[0]
    t = {}
    for i in range(n):
        tag, _, off, ln = struct.unpack('>4sIII', b[12 + 16 * i:28 + 16 * i])
        t[tag.decode('latin-1')] = off
    upm = struct.unpack('>H', b[t['head'] + 18:t['head'] + 20])[0]
    asc, desc = struct.unpack('>hh', b[t['hhea'] + 4:t['hhea'] + 8])
    wa, wd = asc, -desc
    if 'OS/2' in t:
        o = t['OS/2']
        if struct.unpack('>H', b[o + 62:o + 64])[0] & 128:      # USE_TYPO_METRICS
            asc, desc = struct.unpack('>hh', b[o + 68:o + 72])
        wa, wd = struct.unpack('>HH', b[o + 74:o + 78])
    # (A - |D|) / 2 for CSS (hhea/typo) and for the Windows metrics LibreOffice uses in table cells
    return (asc + desc) / 2 / upm, (wa - wd) / 2 / upm


def CAL_DY(it, cell=False):
    """Offset (px) to add to the text frame top so LibreOffice's first baseline lands where Chromium's is.
    Table cells (centred) are laid out by LibreOffice with the font's Windows metrics instead."""
    runs = [r for r in it.get('runs', []) if not r.get('br') and r.get('t', '').strip()]
    if not runs:
        return 0.0
    big = max(runs, key=lambda r: r['size'])
    fs = big['size']
    k, kwin = FONT_K.get(first_family(big.get('family')), (0.335, 0.2635))
    if cell:
        return (k - kwin) * fs
    L = it.get('lh') or fs * 1.4
    return -L / 2 + (k + LO_DESCENT) * fs


_CONVERTED = {}


def ensure_supported_image(path):
    ext = Path(path).suffix.lower()
    if ext in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff'):
        return path
    if path in _CONVERTED:
        return _CONVERTED[path]
    out = Path(WORK) / 'img' / (Path(path).stem + '.png')
    out.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(path) as im:
        im.save(out)
    _CONVERTED[path] = str(out)
    return str(out)


WORK = '.'


def new_presentation(title):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    prs.core_properties.title = title
    prs.core_properties.author = ''
    # theme fonts: Alexandria (headings) / IBM Plex Sans Arabic (body), also for complex script
    for part in prs.part.package.iter_parts():
        if part.partname.endswith('theme1.xml'):
            x = part.blob.decode('utf-8')
            x = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*(")', r'\1Alexandria\2', x)
            x = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*(")', r'\1IBM Plex Sans Arabic\2', x)
            x = re.sub(r'(<a:majorFont>(?:(?!</a:majorFont>).)*?<a:cs typeface=")[^"]*(")', r'\1Alexandria\2', x, flags=re.S)
            x = re.sub(r'(<a:minorFont>(?:(?!</a:minorFont>).)*?<a:cs typeface=")[^"]*(")', r'\1IBM Plex Sans Arabic\2', x, flags=re.S)
            part._blob = x.encode('utf-8')
    return prs


def build_pptx(dom, title, blobs, out_path, diag=False):
    prs = new_presentation(title)
    layout = prs.slide_layouts[6]   # blank
    prev = None
    for d in dom:
        s = prs.slides.add_slide(layout)
        d['enter'] = prev
        SlideBuilder(prs, s, d, diag).build(blobs)
        prev = d.get('transition')
    prs.save(out_path)
    return out_path


# ================================================================= step 4: LibreOffice + QA
def soffice_pdf(pptx_paths, out_dir, work):
    prof = (work / 'lo_profile').resolve()
    pdf_filter = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"},"ExportNotesPages":{"type":"boolean","value":"false"}}'
    cmd = ['soffice', '--headless', '--norestore', '--nologo', f'-env:UserInstallation={prof.as_uri()}',
           '--convert-to', pdf_filter, '--outdir', str(out_dir)] + [str(p) for p in pptx_paths]
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, timeout=1800)


def pdf_to_pngs(pdf, out_dir, prefix, width=None, size=None, aa=True):
    out_dir.mkdir(parents=True, exist_ok=True)
    for f in out_dir.glob(prefix + '-*.png'):
        f.unlink()
    cmd = ['pdftoppm', '-png']
    if size:
        cmd += ['-scale-to-x', str(size[0]), '-scale-to-y', str(size[1])]
    else:
        cmd += ['-scale-to', str(width)]
    if not aa:
        cmd += ['-aa', 'no', '-aaVector', 'no']
    subprocess.run(cmd + [str(pdf), str(out_dir / prefix)], check=True)
    files = sorted(out_dir.glob(prefix + '-*.png'), key=lambda p: int(re.search(r'-(\d+)\.png$', p.name).group(1)))
    return files


def qa_boxes(png):
    """QA colour -> ink bbox (x0, y0, x1, y1) and pixel count."""
    a = np.asarray(Image.open(png).convert('RGB')).astype(np.float32)
    # anti-aliased pixels are blends with the white background; every QA colour has a 0 channel, so
    # alpha = 1 - min/255 and the colour is recovered exactly: c = (p - min) / alpha
    m = a.min(axis=2)
    alpha = 1.0 - m / 255.0
    c = (a - m[..., None]) / np.maximum(alpha, 1e-3)[..., None]
    q = np.rint(c / QA_STEP).astype(np.int16)
    ok = (alpha >= 0.5) & (np.abs(c - q * QA_STEP) <= 9).all(axis=2) & (q.max(axis=2) <= 6) & (q.min(axis=2) >= 0)
    code = q[..., 0] * 49 + q[..., 1] * 7 + q[..., 2]
    ok &= code > 0
    ys, xs = np.nonzero(ok)
    codes = code[ys, xs]
    out = {}
    if len(codes) == 0:
        return out
    order = np.argsort(codes, kind='stable')
    codes, ys, xs = codes[order], ys[order], xs[order]
    cuts = np.flatnonzero(np.diff(codes)) + 1
    for cs, yy, xx in zip(np.split(codes, cuts), np.split(ys, cuts), np.split(xs, cuts)):
        if len(cs) < 3:
            continue
        out[int(cs[0])] = (int(xx.min()), int(yy.min()), int(xx.max()) + 1, int(yy.max()) + 1, len(cs))
    return out


def qa_code(css):
    r, g, b, _ = parse_color(css)
    return (r // QA_STEP) * 49 + (g // QA_STEP) * 7 + (b // QA_STEP)


def plain_text(runs):
    return ''.join(r.get('t', '') for r in collapse_runs(runs) if not r.get('br'))


def text_items(d):
    """(qa colour, box rect, item) for every checked text element / cell of a slide."""
    out = []
    for it in d['items']:
        if it['type'] == 'text' and it.get('qa'):
            out.append((it['qa'], (it['x'], it['y'], it['w'], it['h']), it))
        elif it['type'] == 'table':
            for r in it['rows']:
                for c in r['cells']:
                    if c.get('qa'):
                        out.append((c['qa'], (c['x'], c['y'], c['w'], c['h']), c))
    return out


def check_slide(d, ref_png, lo_png):
    """Compare ink boxes of every text in Chromium vs LibreOffice renders."""
    R, P = qa_boxes(ref_png), qa_boxes(lo_png)
    issues = []
    boxes = []
    for qa, box, it in text_items(d):
        code = qa_code(qa)
        rb, pb = R.get(code), P.get(code)
        txt = plain_text(it['runs'])
        if not txt.strip():
            continue
        lh = it.get('lh') or it.get('fs', 24) * 1.4
        fs = it.get('fs', 24)
        label = txt[:48] + ('...' if len(txt) > 48 else '')
        if rb is None:
            continue
        if pb is None:
            issues.append(('missing', label, 'text not found in the PPTX render'))
            continue
        boxes.append((rb, pb, label))
        dh = (pb[3] - pb[1]) - (rb[3] - rb[1])
        if dh > 0.5 * lh:
            k = max(1, int(round(dh / lh)))
            issues.append(('wrap', label, f'wraps to more lines than the reference (+{k} line{"s" if k > 1 else ""}, {dh:.0f}px taller)'))
        elif dh < -0.6 * lh:
            issues.append(('info', label, f'renders {-dh:.0f}px shorter than the reference (fewer lines: PowerPoint fits more per line)'))
        x, y, w, h = box
        pad = it.get('pad') or [0, 0, 0, 0]
        tol_x = max(3.0, slack_px(it, max(1.0, w - pad[1] - pad[3])) + 2)
        tol_y = max(3.0, 0.3 * lh)
        hx0, hx1 = host_span(it) if it.get('host') else (x, x + w)
        if pb[0] < hx0 - 3 and pb[0] < rb[0] - 3 or pb[2] > hx1 + 3 and pb[2] > rb[2] + 3:
            issues.append(('overflow', label, f'runs past its card/slide edge ({max(hx0 - pb[0], pb[2] - hx1):.0f}px)'))
        def outside(b):
            return (x - b[0], b[2] - (x + w), y - b[1], b[3] - (y + h))
        po, ro = outside(pb), outside(rb)
        ox = max(po[0] - max(ro[0], 0), po[1] - max(ro[1], 0))
        oy = max(po[2] - max(ro[2], 0), po[3] - max(ro[3], 0))
        if ox > tol_x:
            issues.append(('overflow', label, f'overflows its box horizontally by {ox:.0f}px'))
        if oy > tol_y and not dh > 0.5 * lh:
            issues.append(('overflow', label, f'overflows its box vertically by {oy:.0f}px'))
        dy = pb[1] - rb[1]
        if abs(dy) > max(4.0, 0.3 * fs) and abs(dh) <= 0.5 * lh:
            issues.append(('shift', label, f'first line shifted {dy:+.0f}px vertically'))
    # new overlaps between different texts
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            (ri, pi, li), (rj, pj, lj) = boxes[i], boxes[j]
            def inter(a, b):
                return max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(0, min(a[3], b[3]) - max(a[1], b[1]))
            if inter(pi, pj) > 30 and inter(ri, rj) == 0:
                issues.append(('overlap', li, f'overlaps another text: "{lj}"'))
    return issues, R, P


def pdf_line_strings(pdf):
    """Per page: visual-order strings of glyphs per text line (from the PDF glyph positions)."""
    import logging
    import pdfplumber
    logging.getLogger('pdfminer').setLevel(logging.ERROR)
    pages = []
    with pdfplumber.open(str(pdf)) as doc:
        for page in doc.pages:
            chars = [c for c in page.chars if c['text'].strip()]
            lines = []
            for c in sorted(chars, key=lambda c: (round(c['bottom']), c['x0'])):
                for ln in lines:
                    if abs(ln['b'] - c['bottom']) < max(1.5, c['size'] * 0.25):
                        ln['c'].append(c)
                        break
                else:
                    lines.append({'b': c['bottom'], 'c': [c]})
            strs = []
            for ln in lines:
                cs = sorted(ln['c'], key=lambda c: c['x0'])
                strs.append((ln['b'], max(c['size'] for c in cs), ''.join(c['text'] for c in cs)))
            pages.append(strs)
    return pages


def check_tokens(d, page_lines, ref_lines=None):
    """Latin/code tokens inside RTL texts must read left-to-right in the PPTX render (glyph x-order).
    Returns (problems, infos); a token split by a line break into LTR pieces is only info."""
    toks = set()
    for it in d['items']:
        cands = [it] if it['type'] == 'text' else [c for r in it.get('rows', []) for c in r['cells']]
        for c in cands:
            runs = collapse_runs(c['runs'])
            if is_rtl(c['runs'], runs):
                toks.update(ltr_tokens(plain_text(c['runs'])))
    bad, info = [], []
    lines = [(b, sz, l.replace(LRM, '')) for b, sz, l in page_lines]
    blob = '\n'.join(l for _, _, l in lines)
    for t in sorted(toks):
        parts = re.split(r'([-:/.,%])', t)
        rev = ''.join(reversed(parts))
        if t in blob:
            if ref_lines is not None and rev != t:
                rblob = '\n'.join(l for _, _, l in ref_lines)
                if t not in rblob and rev in rblob:
                    info.append((t, f'correct in the PPTX; the Chromium reference shows it reversed as "{rev}" '
                                    '(the source has no LRM marks around it)'))
            continue
        if rev in blob:
            bad.append((t, f'appears reversed ("{rev}")'))
            continue
        wrapped = False
        for k, ch in enumerate(t):
            if ch in '-:/.,%' and 0 < k < len(t) - 1:
                a, b = t[:k + 1], t[k + 1:]
                for b1, sz, l1 in lines:
                    for b2, _, l2 in lines:
                        if 0 < b2 - b1 < 2.5 * sz and a in l1 and b in l2:
                            wrapped = True
        if wrapped:
            info.append((t, 'wrapped at a separator onto the next line (both pieces read left-to-right)'))
        else:
            bad.append((t, 'not found in left-to-right order'))
    return bad, info


def make_sheets(pairs, qa_dir, per_sheet=6):
    """pairs: [(index, id, ref_png, pptx_png)] -> side-by-side sheets + single pairs."""
    pw, ph = 800, 450
    lab = 30
    try:
        font = ImageFont.truetype('DejaVuSans.ttf', 18)
    except Exception:
        font = ImageFont.load_default()
    (qa_dir / 'pairs').mkdir(parents=True, exist_ok=True)
    for f in (qa_dir / 'pairs').glob('*.jpg'):
        f.unlink()
    sheets = []
    for k in range(0, len(pairs), per_sheet):
        chunk = pairs[k:k + per_sheet]
        sheet = Image.new('RGB', (2 * pw + 30, (ph + lab + 10) * len(chunk)), (235, 235, 235))
        dr = ImageDraw.Draw(sheet)
        for i, (idx, sid, ref, lo) in enumerate(chunk):
            y = i * (ph + lab + 10)
            dr.text((10, y + 5), f'{idx:03d}  {sid}   left: reference (Chromium)   right: PPTX rendered by LibreOffice', fill=(20, 20, 20), font=font)
            a = Image.open(ref).convert('RGB').resize((pw, ph), Image.LANCZOS)
            b = Image.open(lo).convert('RGB').resize((pw, ph), Image.LANCZOS)
            sheet.paste(a, (10, y + lab))
            sheet.paste(b, (pw + 20, y + lab))
            pair = Image.new('RGB', (1920 + 10, 540), (235, 235, 235))
            pair.paste(Image.open(ref).convert('RGB').resize((960, 540), Image.LANCZOS), (0, 0))
            pair.paste(Image.open(lo).convert('RGB').resize((960, 540), Image.LANCZOS), (970, 0))
            pair.save(qa_dir / 'pairs' / f'{idx:03d}-{sid}.jpg', quality=88)
        name = qa_dir / f'sheet-{k // per_sheet + 1:02d}.jpg'
        sheet.save(name, quality=85)
        sheets.append(name)
    return sheets


def write_report(path, results, token_bad, dom, pptx, pdf):
    lines = ['# PPTX conversion QA report', '',
             f'- PPTX: `{pptx}`', f'- PDF (LibreOffice render of the PPTX): `{pdf}`',
             f'- Slides: {len(dom)}', '',
             'Method: every text box/table cell is drawn in a unique colour in a QA copy of the slide, once by Chromium '
             '(reference) and once by LibreOffice from a QA copy of the PPTX; their ink bounding boxes are compared: '
             'more lines than the reference, overflow beyond the element box (more than the width slack the converter '
             'grants) or beyond the enclosing card/slide, a shifted first line, new overlaps between texts, missing '
             'text. Latin/code tokens inside RTL text are looked up in left-to-right order in the glyph positions of '
             'the PDF. Non-text objects are checked by eye on the sheets (`qa/sheet-*.jpg`).', '']
    kinds = {}
    for sid, iss in results.items():
        for k, *_ in iss:
            kinds[k] = kinds.get(k, 0) + 1
    info = {s: [x for x in v if x[0] == 'info'] for s, v in results.items()}
    results = {s: [x for x in v if x[0] != 'info'] for s, v in results.items()}
    kinds.pop('info', None)
    ntok = sum(len(v) for v in token_bad.values())
    flagged = sorted(set([s for s, v in results.items() if v] + [s for s, v in token_bad.items() if v]),
                     key=lambda s: [d['id'] for d in dom].index(s))
    lines.append('## Summary')
    lines.append('')
    if not flagged:
        lines.append('No slide has overflowing, clipped, extra-wrapped or reversed text.')
    else:
        lines.append(f'{len(flagged)} slide(s) flagged. Issue counts: ' +
                     ', '.join(f'{k}: {v}' for k, v in sorted(kinds.items())) + (f', reversed/missing tokens: {ntok}' if ntok else ''))
    lines.append('')
    lines.append('| # | slide | issues |')
    lines.append('|---|---|---|')
    idx = {d['id']: d['index'] for d in dom}
    for sid in flagged:
        n = len(results.get(sid, [])) + len(token_bad.get(sid, []))
        lines.append(f'| {idx[sid]} | {sid} | {n} |')
    lines.append('')
    if flagged:
        lines.append('## Details')
        lines.append('')
    for sid in flagged:
        lines.append(f'### {idx[sid]:03d} {sid}')
        for k, label, msg in results.get(sid, []):
            lines.append(f'- **{k}** - "{label}": {msg}')
        for t, msg in token_bad.get(sid, []):
            lines.append(f'- **token** - `{t}`: {msg}')
        lines.append('')
    ninfo = sum(len(v) for v in info.values())
    if ninfo:
        lines.append('## Informational (not problems)')
        lines.append('')
        lines.append('Texts that need fewer lines in the PPTX than in the reference (PowerPoint/LibreOffice fits a '
                     'little more per line, partly because of the width slack; nothing overflows because of it), '
                     'codes that a line break splits at a hyphen, and codes that the Chromium reference shows reversed '
                     'while the PPTX shows them left-to-right.')
        lines.append('')
        for sid in [d['id'] for d in dom]:
            for k, label, msg in info.get(sid, []):
                lines.append(f'- {idx[sid]:03d} {sid} - "{label}": {msg}')
        lines.append('')
    Path(path).write_text('\n'.join(lines), 'utf-8')


# ================================================================= main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--deck', required=True, help='project dir holding deck.json and slides/')
    ap.add_argument('--blobs', required=True, help='JSON map blob id -> local image file')
    ap.add_argument('--out', required=True, help='output .pptx path (PDF, ref/, qa/, fonts/ go next to it)')
    ap.add_argument('--slides', help='comma-separated slide ids (default: deck order)')
    ap.add_argument('--no-qa', action='store_true', help='skip LibreOffice render and QA')
    ap.add_argument('--no-ref-pdf', action='store_true', help='skip the Chromium reference PDF')
    ap.add_argument('--reuse-render', action='store_true', help='reuse _work/dom.json from the last run (needs --keep-work)')
    ap.add_argument('--keep-work', action='store_true', help='keep the intermediate _work folder')
    args = ap.parse_args()

    global WORK
    t0 = time.time()
    deck = Path(args.deck).resolve()
    out = Path(args.out).resolve()
    if out.suffix.lower() != '.pptx':
        out = out.with_suffix('.pptx')
    od = out.parent
    od.mkdir(parents=True, exist_ok=True)
    work = od / '_work'
    work.mkdir(exist_ok=True)
    WORK = str(work)
    ref_dir, qa_dir = od / 'ref', od / 'qa'
    ref_dir.mkdir(exist_ok=True)
    qa_dir.mkdir(exist_ok=True)

    meta = json.loads((deck / 'deck.json').read_text('utf-8'))
    order = meta['order']
    for sid in order:
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', sid):
            raise SystemExit(f'bad slide id in deck.json: {sid!r}')
    if args.slides:
        want = args.slides.split(',')
        order = [s for s in order if s in want]
    blobs = json.loads(Path(args.blobs).read_text('utf-8'))
    title = (meta.get('title') or out.stem).replace(RLE, '')
    for f in (meta.get('faces') or {}).values():
        if f.get('family') not in FONTS:
            log(f'  ! deck face {f.get("family")} is not in the font list; Chromium/LibreOffice will substitute it')

    log('[1/5] fonts + icons')
    files = setup_fonts(od)
    icons = setup_icons()

    if args.reuse_render and (work / 'dom.json').exists():
        log('[2/5] reuse previous Chromium render')
        dom = json.loads((work / 'dom.json').read_text('utf-8'))['slides']
    else:
        log(f'[2/5] Chromium render of {len(order)} slides')
        for f in ref_dir.glob('*.png'):
            f.unlink()
        ref_pdf = None if args.no_ref_pdf else out.with_name(out.stem + '_reference.pdf')
        dom = render(order, deck, blobs, font_css(files), icons, work, ref_dir, ref_pdf)
    for d in dom:
        for w in d.get('warnings', []):
            log(f'  ! {d["id"]}: {w}')

    log('[3/5] build PPTX')
    build_pptx(dom, title, blobs, out)
    log(f'  wrote {out}')
    if args.no_qa:
        if not args.keep_work:
            shutil.rmtree(work, ignore_errors=True)
        return
    diag = work / (out.stem + '_qa.pptx')
    build_pptx(dom, title, blobs, diag, diag=True)

    log('[4/5] LibreOffice -> PDF')
    soffice_pdf([out, diag], od, work)
    pdf = out.with_suffix('.pdf')
    dpdf = od / (diag.stem + '.pdf')
    shutil.move(str(dpdf), str(work / dpdf.name))
    dpdf = work / dpdf.name
    log(f'  wrote {pdf}')

    log('[5/5] QA')
    shutil.rmtree(qa_dir / 'pptx_png', ignore_errors=True)
    lo_pngs = pdf_to_pngs(pdf, qa_dir / 'pptx_png', 'p', width=960)
    diag_pngs = pdf_to_pngs(dpdf, work / 'diag_lo', 'd', size=(SLIDE_W, SLIDE_H), aa=False)
    if len(lo_pngs) != len(dom):
        log(f'  ! page count mismatch: {len(lo_pngs)} pages vs {len(dom)} slides')
    pairs, results, token_bad = [], {}, {}
    page_lines = pdf_line_strings(pdf)
    ref_pdf = out.with_name(out.stem + '_reference.pdf')
    ref_lines = pdf_line_strings(ref_pdf) if ref_pdf.exists() else []
    if len(ref_lines) != len(dom):
        ref_lines = [None] * len(dom)
    for d, lo, dg, pl, rl in zip(dom, lo_pngs, diag_pngs, page_lines, ref_lines):
        i, sid = d['index'], d['id']
        nice = qa_dir / 'pptx_png' / f'{i:03d}-{sid}.png'
        lo.replace(nice)
        ref_png = ref_dir / f'{i:03d}-{sid}.png'
        pairs.append((i, sid, ref_png, nice))
        iss, _, _ = check_slide(d, work / 'diag_ref' / f'{i:03d}-{sid}.png', dg)
        results[sid] = iss
        token_bad[sid], tinfo = check_tokens(d, pl, rl)
        results[sid] += [('info', t, msg) for t, msg in tinfo]
    for f in qa_dir.glob('sheet-*.jpg'):
        f.unlink()
    make_sheets(pairs, qa_dir)
    write_report(qa_dir / 'report.md', results, token_bad, dom, out, pdf)
    nflag = sum(1 for s in results if [x for x in results[s] if x[0] != 'info'] or token_bad[s])
    log(f'  {nflag} slide(s) flagged -> {qa_dir / "report.md"}')
    if not args.keep_work:
        shutil.rmtree(work, ignore_errors=True)
    log(f'done in {time.time() - t0:.0f}s')


if __name__ == '__main__':
    main()
