# Changelog (original slide -> change), source log, unverified items. Run after the final build.
import os, sys, csv, json, importlib
SP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(SP, 'build')); sys.path.insert(0, os.path.join(SP, 'qa')); sys.path.insert(0, os.path.join(SP, 'case_tool'))
os.chdir(os.path.join(SP, 'build'))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from changes import CHANGES
from corrections import ROWS
from common import RDoc
NOT_DONE = json.load(open(os.path.join(SP, 'qa', 'not_done.json'), encoding='utf-8')) if os.path.exists(os.path.join(SP, 'qa', 'not_done.json')) else []
import common  # course-level files: not entity training data
common.FOOTER = common.MARK = 'الكورس 3 · تطبيق الضوابط الأساسية للأمن السيبراني ECC-2:2024'
common.ENTITY_FULL = 'ملف للكورس'

OUT = os.path.join(SP, 'deliver', 'reports')
os.makedirs(OUT, exist_ok=True)
orig = json.load(open(os.path.join(SP, 'qa', 'original_v39', 'project', 'deck.json')))['order']
new = json.load(open('ids.json'))
kinds = {}
for m in ['c_intro', 'c_base', 'c_gov', 'c_def1', 'c_def2', 'c_rtp', 'c_sus', 'c_quiz']:
    for kind, sp in importlib.import_module(m).S:
        kinds[sp['id']] = (kind, sp)
LRM = '‎'
def L(s):  # keep codes LTR in Excel cells
    import re
    return re.sub(r'([0-9A-Za-z][0-9A-Za-z.:/–-]*[-:/][0-9A-Za-z.:/–-]*[0-9A-Za-z])', LRM + r'\1' + LRM, s)

def auto(sid):
    kind, sp = kinds.get(sid, (None, {}))
    notes = []
    if kind in ('evidence',):
        notes.append('شارة نوع المادة، و«يرتبط بالضوابط» بدل «يثبت»، ومربع «لا يثبت وحده».')
    if kind == 'controls':
        fixed = [r[0] for r in sp.get('rows', []) if r[0] in ROWS]
        if fixed:
            notes.append('صياغة «المطلوب» صُححت وفق النص العربي للضوابط: ' + '، '.join(fixed) + '.')
    return ' '.join(notes)

rows = []
for i, sid in enumerate(orig, 1):
    n = new.index(sid) + 1 if sid in new else None
    typ, txt = CHANGES.get(sid, ('', ''))
    extra = auto(sid)
    if not typ and not extra:
        typ, txt = 'دون تغيير في المحتوى', 'تحديث تلقائي فقط: الترقيم، وتثبيت اتجاه الرموز، وشريط التنقل.'
    elif not typ:
        typ = 'تصحيح منهجي'
    if n is None:
        typ, txt = 'نُقل من العرض', 'مفتاح الإجابات نُقل إلى ملف منفصل للمدرب (05_الاختبار_المعرفي).' if sid == 'quiz-ans' else txt
    rows.append([i, sid, n if n else '—', typ, L((txt + ' ' + extra).strip())])
for i, sid in enumerate(new, 1):
    if sid not in orig:
        typ, txt = CHANGES.get(sid, ('شريحة جديدة', ''))
        rows.append(['—', sid, i, typ, L(txt)])
glob_rows = [[k[1:], v[0], L(v[1])] for k, v in CHANGES.items() if k.startswith('*')]

def sheet(wb, name, title, headers, data, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws.append([title]); ws['A1'].font = Font(bold=True, size=13)
    ws.append(headers)
    thin = Side(style='thin', color='9AA5B1')
    for c in ws[2]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9E2F3')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
    for r in data:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if c.row > 2:
                c.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True, readingOrder=2)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'
    ws.page_setup.orientation = 'landscape'  # print all columns on one page width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = '2:2'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{2 + len(data)}'

wb = Workbook(); wb.remove(wb.active)
sheet(wb, 'حسب الشريحة', 'سجل التغييرات: الشريحة الأصلية (الإصدار 39، 174 شريحة) مقابل النسخة المصححة',
      ['رقم الشريحة الأصلي', 'المعرّف', 'رقم الشريحة الجديد', 'نوع التعديل', 'التعديل'], rows, [12, 16, 12, 22, 110])
sheet(wb, 'تعديلات عامة', 'تعديلات طُبقت على مجموعات من الشرائح', ['المعرّف', 'النوع', 'التعديل'], glob_rows, [12, 22, 120])
wb.save(os.path.join(OUT, 'سجل_التغييرات.xlsx'))

# source log
src = list(csv.reader(open(os.path.join(SP, 'verify', 'source_log.csv'), encoding='utf-8-sig')))
extra_src = json.load(open(os.path.join(SP, 'qa', 'extra_sources.json'), encoding='utf-8'))
wb = Workbook(); wb.remove(wb.active)
sheet(wb, 'سجل المصادر', 'سجل المصادر والتحقق', src[0], [[L(x) for x in r] for r in src[1:]] + [[L(x) for x in r] for r in extra_src], [46, 30, 16, 22, 40, 14, 46])
wb.save(os.path.join(OUT, 'سجل_المصادر.xlsx'))

# unverified items
unv = [r for r in src[1:] + extra_src if 'لم يتم التحقق' in r[6]]
IMPACT = {  # how each unverified item is (not) used in the corrected deck
    'مهلة الإبلاغ': 'العرض لا يذكر مهلة زمنية للإبلاغ؛ يكتفي بأن 2-13-3-3 متطلب وطني ويحيل إلى تعليمات الهيئة.',
    'أداة ECC-2:2024': 'العرض يقول إن الهيئة تصدر أداة للتقييم كما ينص الإطار، ولا يذكر اسم الأداة أو إصدارها.',
    'أعداد ومكونات ECC-1': 'العرض يذكر ECC-1:2018 بوصفه الإصدار الأول فقط، دون أعداد أو مقارنة تفصيلية.',
    'وجود مدة 12 شهراً': 'مدة 12 شهراً في العرض منسوبة إلى ECC-2:2024 (2-12-3-5) فقط، لا إلى أطر أخرى.',
    'المعايير الوطنية للتشفير': 'العرض يحيل إلى المعايير الوطنية للتشفير كما يحيل إليها 2-8-3، ولا يسرد خوارزميات أو أطوال مفاتيح.',
    'النص الإنجليزي لـ DCC': 'غير مستخدم؛ الاعتماد على النص العربي لـ DCC-1:2022 (2-6-1-1 إلى 2-6-1-5).',
}
def impact(r):
    for k, v in IMPACT.items():
        if r[0].startswith(k):
            return v
    return 'لا يُنسب في العرض إلى الهيئة؛ يُذكر بوصفه افتراضاً أو يُحال إلى المصدر.'
ASSUME = [
    ['الجهة «هيئة المدار للخدمات الرقمية» وكل بياناتها', 'افتراضية بالكامل: الأسماء والأرقام والتواريخ في نماذج الأدلة والتطبيقات واختبار الأداء من إعداد الكورس.', ['ev-1-9', 'ev-2-2a', 'ev-kpi', 'case-entity', 'perf-e1', 'perf-e5b', 'المرفقات']],
    ['تعريف «الإلغاء المباشر» بساعة واحدة', 'افتراض تدريبي من إجراء الجهة الافتراضية. الضابط 1-9-5 يقول «مباشرة» دون مدة؛ على كل جهة أن تحدد المعنى في إجرائها.', ['ev-1-9', 'ev-1-9b']],
    ['مدد خارطة الطريق (0–3، 3–9، 9–18 شهراً) ومدد التطبيق', 'مدد تخطيط افتراضية، وليست مهلة امتثال من الهيئة؛ الترتيب يتبع المخاطر والاعتماديات والتعليمات النافذة.', ['prog-waves', 'ex7', 'ex7-ans']],
    ['التكرارات في تقويم الاستدامة', 'مقترحة من الكورس، إلا ما نص عليه الإطار أو متطلبات أخرى نافذة على الجهة.', ['review-calendar']],
    ['الأهداف الداخلية للمؤشرات (مثل 90% للسجلات، ونقر 5% أو أقل)', 'أهداف داخلية افتراضية؛ تحقيقها لا يعني اكتمال المتطلب التنظيمي.', ['kpis', 'ev-kpi']],
    ['ممارسات «كيف يُطبق» و«الدليل المقترح» في جداول الضوابط', 'ممارسات مقترحة للتطبيق، وليست متطلبات منسوبة إلى الهيئة؛ عمود «المطلوب» وحده يعيد صياغة النص الرسمي.', ['جداول الضوابط في الوحدات 04 إلى 07']],
]
d = RDoc('نسخة مصححة · ECC-2:2024', title='معلومات لم يتم التحقق منها')
d.title_block('معلومات بقيت دون تحقق', 'قائمة صريحة بما لم يُتحقق منه من المصادر، وسببه، وأثره على العرض', 'ملحق تسليم')
d.table(['المعلومة', 'الإصدار', 'السبب', 'الأثر على العرض'], [[L(r[0]), r[2] or '—', L(r[6].replace('لم يتم التحقق والسبب: ', '')), L(impact(r))] for r in unv], [4.3, 2.7, 5.2, 5.2], size=9.5)
d.h('افتراضات تدريبية (ليست متطلبات من الهيئة)')
where = lambda xs: '، '.join(f'الشريحة {new.index(x) + 1}' if x in new else x for x in xs)
d.table(['الافتراض', 'وضعه', 'أين يظهر (النسخة المصححة)'], [[L(a), L(b), where(c)] for a, b, c in ASSUME], [4.6, 8.0, 4.8], size=9.5)
d.h('فحوص لم تُنفذ')
for x in NOT_DONE:
    d.p('• ' + L(x))
d.save(os.path.join(OUT, 'معلومات_لم_يتم_التحقق_منها.docx'))
print(len(rows), 'changelog rows;', len(src) - 1 + len(extra_src), 'sources;', len(unv), 'unverified')
